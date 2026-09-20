#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
FOTOTAGE MUSTERSTADT 2026 — Werkmeldung übertragen
=============================================
Überträgt die schon ausgefüllten Inhalte aus einer ÄLTEREN Werkmeldung
(z. B. der Datei, die gerade im Umlauf ist) in eine FRISCHE
Vorlage aus build_werkmeldung.py — ohne Datenverlust.

Kopiert werden nur die eingegebenen Werte:
  Blatt „Werke"    — jede ausgefüllte Zeile (Hinweis- und Beispielzeile
                     werden übersprungen), Spalten nach NAME zugeordnet,
                     nicht nach Position. Neue/entfallene Spalten sind egal.
  Blatt „Kontakte" — E-Mail / Telefon / Instagram je Person, nach Name.

Dropdowns, Formeln, Blattschutz, Glossar, Saal Digital, Preishilfe kommen
alle aus der frischen Vorlage. Werte, die jemand in die Preishilfe-Felder
getippt hat, sind persönliche Notizen und werden NICHT übertragen.

Instagram-Handles: Die frische Vorlage bringt die gepflegten Handles aus
data/instagram.csv schon mit. Aus der alten Datei wird ein Handle nur
übernommen, wenn in der Vorlage für die Person KEINER steht. Hat jemand
seinen Handle in der alten Datei korrigiert, gehört die Korrektur in
data/instagram.csv (dann build_werkmeldung.py neu laufen lassen).

Ablauf
------
  1. Neue Vorlage bauen:
       ./.venv/bin/python scripts/build_werkmeldung.py
  2. Die ausgefüllte Datei als Excel (.xlsx) exportieren und in den Ordner
       data/werkmeldung_alt/ legen (eine einzige .xlsx).
  3. Übertragen:
       ./.venv/bin/python scripts/migrate_werkmeldung.py --probe
       ./.venv/bin/python scripts/migrate_werkmeldung.py
  4. Werkmeldung_UEBERTRAGEN.xlsx prüfen, dann als neue
     gemeinsame Datei bereitstellen und die alte ersetzen.

Statt des Ordners geht auch eine Datei direkt:  --alt DEINE_DATEI.xlsx

Wichtig: Während des Umstiegs sollte niemand mehr in der alten Datei
tippen — sonst geht genau diese eine Änderung verloren.
"""

import argparse
import glob
import os
import sys

from openpyxl import load_workbook

import _pfade as P

try:
    from _saal_listen import mal_zeichen
except Exception:                       # pragma: no cover
    def mal_zeichen(t):
        return t

VORLAGE = P.output("Werkmeldung_VORLAGE.xlsx")
ZIEL = P.output("Werkmeldung_UEBERTRAGEN.xlsx")
ORDNER = P.data("werkmeldung_alt")


def finde_alt(ordner):
    """Die eine ausgefüllte .xlsx im Ordner finden (Lock-Dateien ignorieren)."""
    treffer = [p for p in sorted(glob.glob(os.path.join(ordner, "*.xlsx")))
               if not os.path.basename(p).startswith("~$")]
    if not treffer:
        sys.exit(f"Keine .xlsx in {P.rel(ordner)}/ — die ausgefüllte Werkmeldung "
                 f"als Excel-Export (.xlsx) dort ablegen.")
    if len(treffer) > 1:
        sys.exit(f"Mehrere .xlsx in {P.rel(ordner)}/:\n  "
                 + "\n  ".join(os.path.basename(t) for t in treffer)
                 + "\nNur die aktuelle behalten oder mit --alt DATEI.xlsx angeben.")
    return treffer[0]

# Werte, die in der neuen Vorlage umbenannt wurden: alt -> neu.
WERT_MAP = {
    "Papier": {
        "Fujifilm Crystal Archive DP II Silk 232 g/m²":
            "Fujifilm Crystal Archive DP II Silk Portrait 232 g/m²",
    },
}
# Spalten, in denen ein getipptes „x" zum echten „×" gemacht wird.
MASS_SPALTEN = ("Rahmenmaß", "Rahmenmodell")


def _ueberspringen(name, titel):
    n = str(name or "").strip().lower()
    t = str(titel or "").strip().lower()
    return (n.startswith("(") or n.startswith("hinweis")
            or t == "auto eines alten fischers")


def _kopfzeile(ws, suchbegriffe=("Künstler:in", "Titel")):
    """Zeilennummer der Kopfzeile + {Spaltenname: Spaltenindex}."""
    for r in range(1, 15):
        werte = {str(ws.cell(row=r, column=c).value or "").strip(): c
                 for c in range(1, 40)
                 if ws.cell(row=r, column=c).value}
        if all(s in werte for s in suchbegriffe):
            return r, werte
    return None, {}


def lies_werke(pfad):
    """-> [ {Spaltenname: Wert, ...}, ... ]  aus dem Blatt „Werke"."""
    wb = load_workbook(pfad, data_only=True)
    blatt = next((n for n in ("Werke", "Meine Werke") if n in wb.sheetnames), None)
    if not blatt:
        sys.exit(f"FEHLER: Blatt „Werke“ fehlt in {os.path.basename(pfad)}.")
    ws = wb[blatt]
    kopf_r, kopf = _kopfzeile(ws)
    if kopf_r is None:
        sys.exit("FEHLER: Kopfzeile (mit „Künstler:in“ und „Titel“) nicht gefunden.")

    sp_name = kopf["Künstler:in"]
    sp_titel = kopf["Titel"]
    zeilen = []
    for r in range(kopf_r + 1, ws.max_row + 1):
        name = ws.cell(row=r, column=sp_name).value
        titel = ws.cell(row=r, column=sp_titel).value
        if not name and not titel:
            continue
        if _ueberspringen(name, titel):
            continue
        rec = {}
        for spname, ci in kopf.items():
            v = ws.cell(row=r, column=ci).value
            if v is None or (isinstance(v, str) and not v.strip()):
                continue
            rec[spname] = v
        if rec:
            zeilen.append(rec)
    return zeilen


def lies_kontakte(pfad):
    wb = load_workbook(pfad, data_only=True)
    if "Kontakte" not in wb.sheetnames:
        return {}
    ks = wb["Kontakte"]
    _, kopf = _kopfzeile(ks, suchbegriffe=("Künstler:in",))
    if not kopf:
        return {}
    sp_name = kopf["Künstler:in"]

    def hol(row, *titel):
        for t in titel:
            c = kopf.get(t)
            if c and ks.cell(row=row, column=c).value not in (None, ""):
                return ks.cell(row=row, column=c).value
        return None

    out = {}
    for r in range(2, ks.max_row + 1):
        name = ks.cell(row=r, column=sp_name).value
        if not name:
            continue
        out[str(name).strip()] = {
            "E-Mail": hol(r, "E-Mail"),
            "Telefon": hol(r, "Telefon"),
            "Instagram (Handle)": hol(r, "Instagram (Handle)", "Instagram"),
        }
    return out


def fix_wert(spalte, wert):
    wert = WERT_MAP.get(spalte, {}).get(wert, wert)
    if spalte in MASS_SPALTEN:
        wert = mal_zeichen(wert)
    return wert


def main():
    ap = argparse.ArgumentParser(description="ausgefüllte Werkmeldung in die neue Vorlage übertragen")
    ap.add_argument("--ordner", default=ORDNER,
                    help=f"Ordner mit der ausgefüllten .xlsx (Standard: {ORDNER})")
    ap.add_argument("--alt", default=None,
                    help="einzelne .xlsx direkt angeben (statt --ordner)")
    ap.add_argument("--vorlage", default=VORLAGE, help=f"frische Vorlage (Standard: {VORLAGE})")
    ap.add_argument("--ziel", default=ZIEL, help=f"Ausgabedatei (Standard: {ZIEL})")
    ap.add_argument("--probe", action="store_true", help="nur anzeigen, nichts schreiben")
    args = ap.parse_args()

    args.alt = args.alt or finde_alt(args.ordner)
    for p in (args.alt, args.vorlage):
        if not os.path.exists(p):
            sys.exit(f"FEHLER: {p} nicht gefunden.")

    werke = lies_werke(args.alt)
    kontakte = lies_kontakte(args.alt)
    mit_kontakt = sum(1 for v in kontakte.values()
                      if v.get("E-Mail") or v.get("Telefon"))
    print(f"\nAlte Datei: {os.path.basename(args.alt)}")
    print(f"  {len(werke)} ausgefüllte Werkzeile(n)")
    print(f"  {mit_kontakt} Kontakt(e) mit E-Mail/Telefon\n")

    wb = load_workbook(args.vorlage)
    w = wb["Werke"]
    kopf_r, kopf = _kopfzeile(w)
    neue_spalten = set(kopf)
    first = kopf_r + 3            # Kopf, Hinweiszeile, Beispielzeile, dann Eingabe

    uebernommen, ausgelassen = [], set()
    for i, rec in enumerate(werke):
        zr = first + i
        for spname, v in rec.items():
            if spname not in neue_spalten:
                ausgelassen.add(spname)
                continue
            v = fix_wert(spname, v)
            if not args.probe:
                w.cell(row=zr, column=kopf[spname]).value = v
        uebernommen.append(rec.get("Künstler:in", "?") + " — " + str(rec.get("Titel", "?")))

    # Kontakte nach Name zuordnen
    ko = wb["Kontakte"]
    _, kkopf = _kopfzeile(ko, suchbegriffe=("Künstler:in",))
    name_zu_zeile = {}
    for r in range(2, ko.max_row + 1):
        nm = ko.cell(row=r, column=kkopf["Künstler:in"]).value
        if nm:
            name_zu_zeile[str(nm).strip()] = r
    naechste_freie = (max(name_zu_zeile.values()) + 1) if name_zu_zeile else 2

    # Nur SCHREIBEN, wo die neue Vorlage noch leer ist — so bleiben die
    # gepflegten Instagram-Handles der Vorlage erhalten, während die von den
    # Teilnehmenden ergänzten E-Mail/Telefon-Angaben übernommen werden.
    kontakt_erg, kontakt_unbekannt = 0, []
    for name, felder in kontakte.items():
        if not any(felder.values()):
            continue
        r = name_zu_zeile.get(name)
        neu_name = r is None
        if neu_name:
            r = naechste_freie
            naechste_freie += 1
            if not args.probe:
                ko.cell(row=r, column=kkopf["Künstler:in"]).value = name
        geschrieben = False
        for titel, v in felder.items():
            c = kkopf.get(titel) or kkopf.get("Instagram")
            if not (v and c):
                continue
            ziel = ko.cell(row=r, column=c)
            if ziel.value not in (None, "") and not neu_name:
                continue                       # Vorlage hat schon einen Wert
            if not args.probe:
                ziel.value = v
            geschrieben = True
        if geschrieben:
            kontakt_erg += 1
        if neu_name:
            kontakt_unbekannt.append(name)

    # Bericht
    print("Werke, die übertragen werden:")
    for z in uebernommen:
        print("  ·", z)
    if ausgelassen:
        print(f"\nSpalten in der alten Datei ohne Entsprechung in der neuen "
              f"(NICHT übertragen): {', '.join(sorted(ausgelassen))}")
    print(f"\nKontakte ergänzt (E-Mail / Telefon / fehlende Handles): {kontakt_erg}")
    if kontakt_unbekannt:
        print(f"  als NEUE Zeile angelegt (Name nicht in der Vorlage): "
              f"{', '.join(kontakt_unbekannt)}")

    if args.probe:
        print("\nPROBELAUF — nichts geschrieben.\n")
        return

    wb.save(args.ziel)
    print(f"\nGeschrieben: {args.ziel}")
    print("Bitte öffnen und stichprobenartig mit der alten Datei vergleichen, "
          "dann als neue gemeinsame Datei bereitstellen und die alte ersetzen.\n")


if __name__ == "__main__":
    main()
