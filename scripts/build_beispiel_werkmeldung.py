#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
FOTOTAGE MUSTERSTADT 2026 — Beispiel-Werkmeldung
================================================
Erzeugt  data/beispiele/Werkmeldung_BEISPIEL.xlsx : eine schon AUSGEFÜLLTE
Werkmeldung mit frei erfundenen Werken der Beispiel-Teilnehmenden aus
data/teilnehmende.txt. Sie dient zum Ausprobieren der Skripte, die sonst
eine echte, ausgefüllte Werkmeldung brauchen:

  Einlesen in den Master (ohne die echte Ablage data/werkmeldungen/ anzufassen):
      ./.venv/bin/python scripts/import_werkmeldungen.py --ordner data/beispiele --probe

  Übertragen in eine frische Vorlage:
      ./.venv/bin/python scripts/build_werkmeldung.py
      ./.venv/bin/python scripts/migrate_werkmeldung.py --alt data/beispiele/Werkmeldung_BEISPIEL.xlsx --probe

Die Datei liegt bewusst NICHT in data/werkmeldungen/ oder data/werkmeldung_alt/:
dort würde sie beim nächsten Einlesen als echte Meldung behandelt.
"""

import os

from openpyxl import load_workbook

import _pfade as P
import _teilnehmende as T
import build_werkmeldung as B

ZIEL = P.data("beispiele", "Werkmeldung_BEISPIEL.xlsx")

# Reihenfolge wie die Spalten in build_werkmeldung.SPALTEN:
# Künstler:in, Titel, Ort, Land, Jahr, Druckverfahren, Papier,
# Bildmaß H, Bildmaß B, Blattmaß H, Blattmaß B, Auflage, AP, Exemplar,
# Rahmenmodell, Rahmenmaß, Passepartout, Preis, Rahmen im Preis,
# Versicherungswert, Verkäuflich
FOTO = "Fujifilm Crystal Archive DP II Silk Portrait 232 g/m²"
RAG = "Hahnemühle Photo Rag 308 g/m²"
BARYTA = "Hahnemühle FineArt Baryta 325 g/m²"
RAHMEN_40 = "Beispielrahmen 40x50 cm, schwarz"
RAHMEN_50 = "Beispielrahmen 50x70 cm, alu roh"

WERKE = [
    ("Mia Beispiel", "Morgen am Hafen", "Musterstadt", "Deutschland", 2024, "C-Print", FOTO,
     27, 40, 30, 45, 5, 1, "1/5", RAHMEN_40, "40x50 cm", "weiß, 2 mm", 90, "ja", 140, "ja"),
    ("Mia Beispiel", "Netze", "Musterstadt", "Deutschland", 2024, "C-Print", FOTO,
     40, 27, 45, 30, 5, 1, "2/5", RAHMEN_40, "40x50 cm", "weiß, 2 mm", 90, "ja", 140, "ja"),
    ("Noah Muster", "Zwischenraum I", "Beispieldorf", "Deutschland", 2025, "Pigmentdruck", RAG,
     45, 60, 50, 65, 3, 1, "1/3", RAHMEN_50, "50x70 cm", "ohne", 240, "ja", 380, "ja"),
    ("Noah Muster", "Zwischenraum II", "Beispieldorf", "Deutschland", 2025, "Pigmentdruck", RAG,
     45, 60, 50, 65, 3, 1, "2/3", RAHMEN_50, "50x70 cm", "ohne", 240, "ja", 380, "ja"),
    ("Lina Vogel", "Ohne Titel (Küche)", "Beispielstadt", "Österreich", 2023,
     "Silbergelatine, Handabzug", "anderes Papier — bitte in die Zelle schreiben",
     24, 30, 30, 40, 8, 2, "3/8", "Magnetrahmen 30x40 cm", "30x40 cm", "weiß, 2 mm",
     140, "ja", 200, "ja"),
    ("Lina Vogel", "Ohne Titel (Treppe)", "Beispielstadt", "Österreich", 2023,
     "Silbergelatine, Handabzug", "anderes Papier — bitte in die Zelle schreiben",
     30, 24, 40, 30, 8, 2, "1/8", "Magnetrahmen 30x40 cm", "30x40 cm", "weiß, 2 mm",
     None, None, 200, "nein"),
    ("Paul Winter", "Nachtschicht", "Musterhausen", "Deutschland", 2025, "Pigmentdruck", BARYTA,
     50, 75, 55, 80, 5, 1, "1/5", "Eichenrahmen 55x80 cm", "55x80 cm", "ohne",
     420, "ja", 600, "ja"),
    ("Sara Klein", "Grenzverlauf", "Beispielgrenze", "Luxemburg", 2026, "Pigmentdruck", RAG,
     30, 30, 40, 40, 3, 1, "1/3", "Beispielrahmen 40x40 cm, weiß", "40x40 cm", "ohne",
     160, "nein", 240, "ja"),
    ("Sara Klein", "Grenzverlauf II", "Beispielgrenze", "Luxemburg", 2026, "Pigmentdruck", RAG,
     30, 30, 40, 40, 3, 1, "2/3", "Beispielrahmen 40x40 cm, weiß", "40x40 cm", "ohne",
     160, "nein", 240, "ja"),
    ("Tom Feldmann", "Blau, sonntags", "Musterstadt", "Deutschland", 2024, "Cyanotypie",
     "anderes Papier — bitte in die Zelle schreiben",
     20, 30, 25, 35, 1, 0, "1/1", "Bilderleiste 25x35 cm", "25x35 cm", "ohne",
     180, "ja", 180, "ja"),
]

# E-Mail / Telefon je Person (Instagram kommt schon aus data/instagram.csv)
KONTAKTE = {
    "Mia Beispiel": ("mia.beispiel@example.org", "+49 170 0000101"),
    "Noah Muster": ("noah.muster@example.org", "+49 170 0000102"),
    "Lina Vogel": ("lina.vogel@example.org", "+49 170 0000103"),
    "Paul Winter": ("paul.winter@example.org", "+49 170 0000104"),
    "Sara Klein": ("sara.klein@example.org", "+49 170 0000105"),
    "Tom Feldmann": ("tom.feldmann@example.org", "+49 170 0000106"),
}


def main():
    fehlt = sorted({w[0] for w in WERKE} - set(T.NAMEN))
    if fehlt:
        raise SystemExit("Diese Namen fehlen in data/teilnehmende.txt: " + ", ".join(fehlt))

    os.makedirs(os.path.dirname(ZIEL), exist_ok=True)
    B.OUT = ZIEL            # Vorlage direkt an den Zielort bauen, output/ bleibt unberührt
    B.build()

    wb = load_workbook(ZIEL)
    ws = wb["Werke"]
    erste = 4               # Kopf, Hinweiszeile, Beispielzeile, dann die Eingabe
    for r, werk in enumerate(WERKE, start=erste):
        for c, wert in enumerate(werk, start=1):
            ws.cell(row=r, column=c, value=wert)

    ko = wb["Kontakte"]
    for r in range(2, ko.max_row + 1):
        name = ko.cell(row=r, column=1).value
        if name in KONTAKTE:
            ko.cell(row=r, column=2, value=KONTAKTE[name][0])
            ko.cell(row=r, column=3, value=KONTAKTE[name][1])

    wb.save(ZIEL)
    print(f"geschrieben: {P.rel(ZIEL)}  ({len(WERKE)} Werke, {len(KONTAKTE)} Kontakte)")


if __name__ == "__main__":
    main()
