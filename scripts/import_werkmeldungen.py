#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
FOTOTAGE MUSTERSTADT 2026 — Werkmeldungen einsammeln
================================================
Liest die gemeinsame Werkmeldung (ein Blatt „Werke" mit einer Spalte
„Künstler:in" pro Zeile, Kontaktdaten im Blatt „Kontakte") und hängt alle
Werke an das Blatt „Werke" des Masters an. Legt fehlende Teilnehmende im
Blatt „Teilnehmende" an.

Werknummern werden NICHT in Meldereihenfolge vergeben, sondern nach jedem
Import komplett neu durchnummeriert — sortiert nach Künstler:in (A → Z),
bei gleichem Namen nach Titel. Das betrifft alle „echten" Werke im Master,
nicht nur die neu importierten, und auch nicht die Beispielwerke (Spalte
„Beispiel" = ja). Dadurch ist die Nummer eines Werks nicht stabil, solange
noch Meldungen fehlen — erst nach dem letzten Import vor dem Druck sind
die Nummern final. Bereits gedruckte Wandschilder mit einer älteren Nummer
stimmen danach nicht mehr überein.

Es können auch mehrere Dateien im Ordner liegen (z. B. eine pro
Google-Drive-Export) — sie werden alle gelesen.

KEINE Dublettenprüfung: jede Meldungszeile wird als eigenes Werk angehängt,
auch wenn Künstler:in, Titel, Maße, Rahmen und Preis exakt einer bereits
importierten Zeile gleichen. Dieselbe Datei zweimal einlesen verdoppelt
also alle Werke — vor einem erneuten Lauf mit derselben Quelle den Master
auf den Beispiel-Stand zurücksetzen (Sicherung + Zeilen ab „echte Werke"
löschen) statt einfach neu zu importieren.

Aufruf
------
  ./.venv/bin/python scripts/import_werkmeldungen.py
  ./.venv/bin/python scripts/import_werkmeldungen.py --probe   # nur anzeigen
  (oder über  scripts/fub.py , Menüpunkt 3)

Ablage der ausgefüllten Werkmeldung(en): data/werkmeldungen/
Der Master wird vor jedem Schreibvorgang automatisch gesichert als
output/FUB2026_Werkdaten_MASTER_backup_JJJJMMTT-HHMM.xlsx
"""

import argparse
import glob
import os
import shutil
import sys
from datetime import datetime

from openpyxl import load_workbook
from openpyxl.styles import Font, Border, Side

import _pfade as P
import _saal_listen as SL

MASTER = P.output("FUB2026_Werkdaten_MASTER.xlsx")
thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

# Werkmeldung -> Master
MAPPING = {
    "Titel": "Titel",
    "Ort": "Ort",
    "Land": "Land",
    "Jahr": "Jahr",
    "Druckverfahren": "Druckverfahren",
    "Papier": "Papier",
    "Bildmaß H cm": "Bildmaß H cm",
    "Bildmaß B cm": "Bildmaß B cm",
    "Blattmaß H cm": "Blattmaß H cm",
    "Blattmaß B cm": "Blattmaß B cm",
    "Auflage": "Auflage",
    "AP": "AP",
    "Exemplar": "Exemplar",
    "Rahmenmodell": "Rahmenmodell",
    "Rahmenmaß": "Rahmenmaß",
    "Passepartout": "Passepartout",
    "Preis EUR": "Preis EUR",
    "Rahmen im Preis": "Rahmen im Preis",
    "Versicherungswert EUR": "Versicherungswert EUR",
}
BEISPIELMARKER = ("(beispiel", "auto eines alten fischers")

def _ueberspringen(name, titel):
    """Hinweis- und Beispielzeile der Vorlage erkennen."""
    n = str(name or "").strip().lower()
    t = str(titel or "").strip().lower()
    return (n.startswith("(") or n.startswith("hinweis")
            or t == "auto eines alten fischers")


def lies_kontakte(wb):
    """{Name: {'E-Mail':.., 'Telefon':.., 'Instagram':..}} aus Blatt „Kontakte"."""
    out = {}
    if "Kontakte" not in wb.sheetnames:
        return out
    ks = wb["Kontakte"]
    kopf = {str(c.value).strip(): c.column for c in ks[1] if c.value}
    sp_name = kopf.get("Künstler:in", 1)
    for r in range(2, ks.max_row + 1):
        name = ks.cell(row=r, column=sp_name).value
        if not name:
            continue
        def g(titel):
            c = kopf.get(titel)
            return ks.cell(row=r, column=c).value if c else None
        out[str(name).strip()] = {
            "E-Mail": g("E-Mail"),
            "Telefon": g("Telefon"),
            "Instagram": g("Instagram (Handle)") or g("Instagram"),
        }
    return out


def lies_meldung(pfad):
    """-> (kontakte_dict, [werk_dict, ...])"""
    wb = load_workbook(pfad, data_only=True)
    blatt = "Werke" if "Werke" in wb.sheetnames else (
        "Meine Werke" if "Meine Werke" in wb.sheetnames else None)
    if not blatt:
        print(f"  übersprungen (Blatt „Werke“ fehlt): {os.path.basename(pfad)}")
        return {}, []
    ws = wb[blatt]
    kontakte = lies_kontakte(wb)

    header_row = None
    for r in range(1, 12):
        werte = [str(ws.cell(row=r, column=c).value or "").strip() for c in range(1, 30)]
        if "Künstler:in" in werte and "Titel" in werte:
            header_row = r
            break
    if header_row is None:
        print(f"  übersprungen (Kopfzeile nicht gefunden): {os.path.basename(pfad)}")
        return {}, []

    header = {}
    for c in range(1, 40):
        v = ws.cell(row=header_row, column=c).value
        if v:
            header[str(v).strip()] = c

    sp_name = header["Künstler:in"]
    sp_titel = header["Titel"]
    werke = []
    for r in range(header_row + 1, ws.max_row + 1):
        name = ws.cell(row=r, column=sp_name).value
        titel = ws.cell(row=r, column=sp_titel).value
        if not name and not titel:
            continue
        if _ueberspringen(name, titel):
            continue
        if not name:
            print(f"    Zeile {r}: „{titel}“ ohne Künstler:in — übersprungen")
            continue
        if not titel:
            print(f"    Zeile {r}: {name} ohne Titel — übersprungen")
            continue
        w = {"Künstler:in": str(name).strip()}
        for quelle, ziel in MAPPING.items():
            if quelle in header:
                v = ws.cell(row=r, column=header[quelle]).value
                w[ziel] = v.strip() if isinstance(v, str) else v   # Rand-Leerzeichen weg
        # „40x50" -> „40×50": echtes Malzeichen in den Maß-Textfeldern
        for feld in ("Rahmenmaß", "Rahmenmodell"):
            if w.get(feld):
                w[feld] = SL.mal_zeichen(w[feld])
        verk = ws.cell(row=r, column=header["Verkäuflich"]).value if "Verkäuflich" in header else "ja"
        w["Status"] = "verfügbar" if str(verk or "ja").strip().lower() == "ja" else "nicht verkäuflich"
        if w["Status"] == "nicht verkäuflich":
            w["Preis EUR"] = None
            w["Rahmen im Preis"] = None
        werke.append(w)
    return kontakte, werke


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--ordner", default=P.data("werkmeldungen"))
    ap.add_argument("--master", default=MASTER)
    ap.add_argument("--probe", action="store_true", help="nur anzeigen, nichts schreiben")
    args = ap.parse_args()

    if not os.path.exists(args.master):
        sys.exit(f"FEHLER: {args.master} nicht gefunden.")
    dateien = sorted(glob.glob(os.path.join(args.ordner, "*.xlsx")))
    dateien = [d for d in dateien if not os.path.basename(d).startswith("~$")]
    if not dateien:
        sys.exit(f"Keine .xlsx-Dateien in {P.rel(args.ordner)}/")

    print(f"\n{len(dateien)} Datei(en) gefunden.\n")

    wb = load_workbook(args.master)
    ws = wb["Werke"]
    wt = wb["Teilnehmende"]
    header = {c.value: c.column for c in ws[1] if c.value}

    namen_mit_werk = set()
    erste_leere = None
    for r in range(2, ws.max_row + 1):
        t = ws.cell(row=r, column=header["Titel"]).value
        k = ws.cell(row=r, column=header["Künstler:in"]).value
        if t or k:
            if k:
                namen_mit_werk.add(str(k).strip())
        elif erste_leere is None:
            erste_leere = r
    if erste_leere is None:
        erste_leere = ws.max_row + 1

    teilnehmende = {}
    for r in range(2, wt.max_row + 1):
        n = wt.cell(row=r, column=1).value
        if n:
            teilnehmende[str(n).strip()] = r

    kontakte_gesamt = {}
    zeile = erste_leere
    neu_gesamt = 0
    for d in dateien:
        print(f"» {os.path.basename(d)}")
        kontakte, werke = lies_meldung(d)
        kontakte_gesamt.update({k: v for k, v in kontakte.items() if any(v.values())})
        pro_person = {}
        for w in werke:
            # Keine Dublettenprüfung: jede Meldungszeile ist ein eigenes Werk,
            # unabhängig davon, ob Titel/Maße/Rahmen einer anderen Zeile gleichen.
            name = w["Künstler:in"]
            if not args.probe:
                for feld, wert in w.items():
                    if feld in header:
                        c = ws.cell(row=zeile, column=header[feld], value=wert)
                        c.font = Font(name="Arial", size=10, color="0000FF")
                        c.border = BORDER
            namen_mit_werk.add(name)
            zeile += 1
            neu_gesamt += 1
            pro_person[name] = pro_person.get(name, 0) + 1
        for name, n in sorted(pro_person.items()):
            print(f"    {name}: {n} neue Werke")

    # Werknummern automatisch vergeben: alle "echten" Werke (Spalte "Beispiel"
    # ungleich "ja") komplett neu durchnummeriert, sortiert nach Künstler:in
    # (A-Z), bei gleichem Namen nach Titel — unabhängig davon, in welcher
    # Zeile bzw. Reihenfolge sie gemeldet/importiert wurden. Beispielwerke
    # bleiben unangetastet, damit sie im Testlauf weiter 001-014 zeigen.
    sp_beispiel = header.get("Beispiel")
    echte_zeilen = []
    for r in range(2, ws.max_row + 1):
        k = ws.cell(row=r, column=header["Künstler:in"]).value
        t = ws.cell(row=r, column=header["Titel"]).value
        if not k and not t:
            continue
        bsp = ws.cell(row=r, column=sp_beispiel).value if sp_beispiel else None
        if str(bsp or "").strip().lower() in ("ja", "yes", "true", "1", "x"):
            continue
        echte_zeilen.append((str(k or "").strip(), str(t or "").strip(), r))
    echte_zeilen.sort(key=lambda z: (z[0].casefold(), z[1].casefold()))

    neu_nummeriert = 0
    for i, (k, t, r) in enumerate(echte_zeilen, start=1):
        neue_nr = f"{i:03d}"
        c = ws.cell(row=r, column=header["Nr"])
        if str(c.value or "") != neue_nr:
            neu_nummeriert += 1
        if not args.probe:
            c.value = neue_nr
            c.number_format = "@"
    if echte_zeilen:
        print(f"\nWerknummern nach Namen (A-Z) vergeben: 001–{len(echte_zeilen):03d} "
              f"({neu_nummeriert} davon geändert)"
              + (" — Probelauf, nicht geschrieben" if args.probe else "") + ".")

    # Teilnehmende anlegen bzw. leere Kontaktfelder ergänzen
    for name in sorted(namen_mit_werk):
        ktk = kontakte_gesamt.get(name, {})
        werte = {2: ktk.get("E-Mail"), 3: ktk.get("Telefon"), 4: ktk.get("Instagram")}
        if name in teilnehmende:
            r = teilnehmende[name]
            luecken = [i for i, v in werte.items()
                       if v and not wt.cell(row=r, column=i).value]
            if luecken and not args.probe:
                for i in luecken:
                    c = wt.cell(row=r, column=i, value=werte[i])
                    c.font = Font(name="Arial", size=10, color="0000FF")
                    c.border = BORDER
                print(f"    Kontakt ergänzt: {name}")
            continue
        if args.probe:
            print(f"    Teilnehmende:in neu: {name}")
            continue
        r = wt.max_row + 1
        for rr in range(2, wt.max_row + 2):
            if not wt.cell(row=rr, column=1).value:
                r = rr
                break
        for i, v in enumerate([name, werte[2], werte[3], werte[4]], start=1):
            c = wt.cell(row=r, column=i, value=v)
            c.font = Font(name="Arial", size=10, color="0000FF")
            c.border = BORDER
        teilnehmende[name] = r
        print(f"    Teilnehmende:in neu angelegt: {name}")

    if args.probe:
        print(f"\nPROBELAUF — nichts geschrieben. {neu_gesamt} Werke würden importiert.")
        print()
        return

    if neu_gesamt or neu_nummeriert:
        stamp = datetime.now().strftime("%Y%m%d-%H%M")
        backup = args.master.replace(".xlsx", f"_backup_{stamp}.xlsx")
        shutil.copy(args.master, backup)
        wb.save(args.master)
        print(f"\n{neu_gesamt} Werke importiert, {neu_nummeriert} Werknummer(n) angepasst. "
              f"Sicherung: {backup}")
        print("ACHTUNG: Die Formelspalten 'Maßangabe' und 'Auflageangabe' berechnet Excel "
              "beim nächsten Öffnen. Der Generator braucht sie nicht — er rechnet selbst.\n")
    else:
        print("\nNichts Neues zu importieren.\n")


if __name__ == "__main__":
    main()
