#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
FOTOTAGE MUSTERSTADT 2026 — Druckdaten-Generator (Konzept A)
======================================================
Liest output/Werkdaten_MASTER.xlsx und erzeugt in output/ausgabe/:

  01_Wandschilder.pdf         DIN A0–A10 quer (Standard A8, 10/A4); Schnittlinien + QR
  02_Werkliste.pdf            ausliegende Werk- und Preisliste, A4, 2-spaltig, fortlaufend
  03_Rueckseitenetiketten.pdf Etiketten für die Rahmenrückseite (12/A4)
  04_Versicherungsliste.pdf   interne Liste mit Versicherungswerten
  05_Instagram_Wand.pdf       Aushang mit den Instagram-QR-Codes aller Teilnehmenden
  Werkdaten_Serienbrief.csv   Fallback für Word-Serienbrief

Aufruf
------
  ./.venv/bin/python scripts/build_druckdaten.py
  ./.venv/bin/python scripts/build_druckdaten.py --schildformat a7      # größer, 8/A4
  ./.venv/bin/python scripts/build_druckdaten.py --schildformat a6      # groß/barrierearm, 4/A4
  ./.venv/bin/python scripts/build_druckdaten.py --schildformat a3      # jedes DIN-A-Format von a0 bis a10
  ./.venv/bin/python scripts/build_druckdaten.py --einzelkarte 070 --schildformat a2   # ein Schild pro Blatt
  ./.venv/bin/python scripts/build_druckdaten.py --nur wandschilder
  ./.venv/bin/python scripts/build_druckdaten.py --gruppierung Wand     # Werkliste nach Wänden
  ./.venv/bin/python scripts/build_druckdaten.py --ohne-preishinweis
  (oder bequem über  scripts/menue.py , Menüpunkt 2)

Die 14 Beispielwerke aus build_master.py (Spalte "Beispiel" = ja) werden automatisch
übersprungen, sobald mindestens ein echtes Werk im Master steht. --mit-beispielen
erzwingt sie trotzdem mit (z. B. für einen Layout-Test mit vollem Datensatz).

Alle Ausgaben folgen der Reihenfolge der Spalte "Nr" — die vergibt
import_werkmeldungen.py beim Einsammeln der Meldungen automatisch alphabetisch
nach Künstler:in (A-Z), nicht in Melde- oder Zeilenreihenfolge.

Abhängigkeiten: reportlab, openpyxl, segno   (pip install reportlab openpyxl segno)
"""

import argparse
import csv
import json
import os
import re
import sys
from datetime import date
from io import BytesIO

from openpyxl import load_workbook
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.units import mm
from reportlab.lib.utils import ImageReader
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas

import _pfade as P

try:
    import _teilnehmende as T
except Exception:                       # pragma: no cover
    T = None

try:
    from _saal_listen import mal_zeichen
except Exception:                       # pragma: no cover
    import re as _re
    def mal_zeichen(text):
        """„40x50" -> „40×50" (Fallback, falls _saal_listen.py fehlt)."""
        return text if text is None else _re.sub(r"(\d\s*)[xX](\s*\d)", r"\1×\2", str(text))

# ----------------------------------------------------------------- Ausstellung
AUSSTELLUNG = {
    "titel": "FOTOTAGE MUSTERSTADT 2026",
    "untertitel": "Ausstellung für zeitgenössische Fotografie",
    "ort": "Kunstraum Musterstadt",
    "laufzeit": "12.—19. September 2026",
    "web": "fototage-musterstadt.de",
    "mail": "kontakt@beispiel-verein.de",
    "insta": "instagram.com/fototage.musterstadt",
}

# ----------------------------------------------------------------- Schriften
# Eigene Schrift: einfach Regular- und Bold-Schnitt (.ttf/.otf, keine Variable
# Fonts) in SCHRIFTEN_ORDNER ablegen — keine festen Dateinamen, keine Pfade in
# design.json nötig. Erkennung per Dateiname: "bold"/"fett" = Bold, alles
# andere ohne "italic"/"oblique"/"kursiv" = Regular. Es werden bewusst nur
# diese zwei Schnitte verwendet: italic/bolditalic zeigen im PDF auf dieselben
# Dateien wie regular/bold, damit nirgends echter Kursivtext gerendert wird.
# Gilt NUR für die PDFs hier (ausgabe/) — die xlsx-Dateien (build_master.py,
# build_werkmeldung.py) tragen nur den Font-NAMEN "Arial" ein und sind davon
# unabhängig.
SCHRIFTEN_ORDNER = P.assets("schriften")
_SCHRIFT_ENDUNGEN = (".ttf", ".otf")
_ITALIC_WOERTER = ("italic", "oblique", "kursiv")
_BOLD_WOERTER = ("bold", "fett")

FONT_DIRS = [
    "/usr/share/fonts/truetype/liberation",
    "/usr/share/fonts/truetype/dejavu",
    "C:/Windows/Fonts",
    "/System/Library/Fonts/Supplemental",
    "/Library/Fonts",
]
FONT_SETS = [
    ("LiberationSans-Regular.ttf", "LiberationSans-Bold.ttf",
     "LiberationSans-Italic.ttf", "LiberationSans-BoldItalic.ttf"),
    ("Arial.ttf", "Arial Bold.ttf", "Arial Italic.ttf", "Arial Bold Italic.ttf"),
    ("arial.ttf", "arialbd.ttf", "ariali.ttf", "arialbi.ttf"),
    ("DejaVuSans.ttf", "DejaVuSans-Bold.ttf",
     "DejaVuSans-Oblique.ttf", "DejaVuSans-BoldOblique.ttf"),
]

F_REG, F_BOLD, F_IT, F_BI = "Helvetica", "Helvetica-Bold", "Helvetica-Oblique", "Helvetica-BoldOblique"


def finde_schriften(ordner):
    """Sucht Regular- und Bold-Schnitt im Ordner, per Dateiname erkannt (siehe
    oben). -> {'regular':pfad,'bold':pfad,'italic':pfad,'bolditalic':pfad}
    (italic/bolditalic = Aliase auf regular/bold) oder None, falls nicht
    beide gefunden wurden."""
    if not os.path.isdir(ordner):
        return None
    dateien = sorted(f for f in os.listdir(ordner) if f.lower().endswith(_SCHRIFT_ENDUNGEN))

    def ist(f, woerter):
        n = f.lower()
        return any(w in n for w in woerter)

    bold = next((f for f in dateien if ist(f, _BOLD_WOERTER) and not ist(f, _ITALIC_WOERTER)), None)
    regular = next((f for f in dateien if not ist(f, _BOLD_WOERTER) and not ist(f, _ITALIC_WOERTER)), None)
    if not regular or not bold:
        return None
    r, b = os.path.join(ordner, regular), os.path.join(ordner, bold)
    return {"regular": r, "bold": b, "italic": r, "bolditalic": b}


def register_fonts():
    """Sucht zuerst eigene Schriften in SCHRIFTEN_ORDNER, dann eine
    Arial-metrische Systemschrift; fällt sonst auf Helvetica zurück."""
    global F_REG, F_BOLD, F_IT, F_BI
    eigene = finde_schriften(SCHRIFTEN_ORDNER)
    if eigene:
        try:
            pdfmetrics.registerFont(TTFont("SCHRIFT", eigene["regular"]))
            pdfmetrics.registerFont(TTFont("SCHRIFT-B", eigene["bold"]))
            pdfmetrics.registerFont(TTFont("SCHRIFT-I", eigene["italic"]))
            pdfmetrics.registerFont(TTFont("SCHRIFT-BI", eigene["bolditalic"]))
            F_REG, F_BOLD, F_IT, F_BI = "SCHRIFT", "SCHRIFT-B", "SCHRIFT-I", "SCHRIFT-BI"
            return f"{os.path.basename(eigene['regular'])} + {os.path.basename(eigene['bold'])} (aus {SCHRIFTEN_ORDNER}/)"
        except Exception:
            pass
    for names in FONT_SETS:
        for d in FONT_DIRS:
            paths = [os.path.join(d, n) for n in names]
            if all(os.path.exists(p) for p in paths):
                try:
                    pdfmetrics.registerFont(TTFont("SCHRIFT", paths[0]))
                    pdfmetrics.registerFont(TTFont("SCHRIFT-B", paths[1]))
                    pdfmetrics.registerFont(TTFont("SCHRIFT-I", paths[2]))
                    pdfmetrics.registerFont(TTFont("SCHRIFT-BI", paths[3]))
                    F_REG, F_BOLD, F_IT, F_BI = "SCHRIFT", "SCHRIFT-B", "SCHRIFT-I", "SCHRIFT-BI"
                    return os.path.basename(paths[0])
                except Exception:
                    continue
    return "Helvetica (Standard)"


# ----------------------------------------------------------------- Design-Konfiguration
# Überschreibungen aus design.json ("wandschild"-Block). Wirken in schild_cfg().
DESIGN_SCHILD = {}          # {"a7": {...}, "a5": {...}}  — je Schildformat
ANPASSUNG = {               # automatische Anpassung der Inhalte, siehe data/README_design.md
    "aktiv": True,
    "text_min_faktor": 0.7,      # einzelne Zeile darf bis auf 70 % schrumpfen, bevor sie gekürzt wird
    "inhalt_min_faktor": 0.4,    # gesamtes Schild darf bis auf 40 % der Sollgrößen schrumpfen
    "bogenrand_mm": 8.0,         # Rand für Schnittmarken beim Anordnen mehrerer Schilder pro Bogen
}


def lade_design(pfad):
    """Liest design.json des Designers und merkt sich Schildmaße, Typografie-
    größen und Farben. Die Schrift selbst kommt unabhängig davon aus
    register_fonts() (Ordner schriften/, automatisch erkannt). Die Schild-
    Überschreibungen werden erst in schild_cfg() angewendet."""
    if not pfad:
        return None
    if not os.path.exists(pfad):
        sys.exit(f"FEHLER: Designdatei {pfad} nicht gefunden.")
    with open(pfad, encoding="utf-8") as f:
        d = json.load(f)

    ws = d.get("wandschild") or {}
    for fmt, blk in ws.items():
        if fmt.startswith("_"):
            continue
        if fmt == "anpassung":
            for k in ANPASSUNG:
                if k in blk:
                    ANPASSUNG[k] = blk[k]
            continue
        if fmt.lower() not in FORMATE:
            print(f"Hinweis: design.json → wandschild → „{fmt}“ ist kein DIN-A-Format "
                  f"(erlaubt: {', '.join(FORMATE)}) — ignoriert.")
            continue
        if isinstance(blk, dict):
            DESIGN_SCHILD[fmt.lower()] = blk
    return d


def hex_rgb(h, fallback=(0, 0, 0)):
    try:
        h = h.lstrip("#")
        return tuple(int(h[i:i + 2], 16) / 255 for i in (0, 2, 4))
    except Exception:
        return fallback


FARBEN = {"text": (0, 0, 0), "sekundaer": (0.38, 0.38, 0.38),
          "tertiaer": (0.60, 0.60, 0.60), "linie": (0.78, 0.78, 0.78),
          "signal": (0.70, 0.12, 0.12)}


def setze_farben(d):
    if not d:
        return
    for k, v in (d.get("farben") or {}).items():
        if k in FARBEN:
            FARBEN[k] = hex_rgb(v, FARBEN[k])


# ----------------------------------------------------------------- Daten lesen
SPALTEN = [
    "Nr", "Wand", "Position", "Künstler:in", "Titel", "Ort", "Land", "Jahr",
    "Druckverfahren", "Papier", "Bildmaß H cm", "Bildmaß B cm", "Blattmaß H cm",
    "Blattmaß B cm", "Maßangabe", "Auflage", "AP", "Exemplar", "Auflageangabe",
    "Rahmenmodell", "Rahmenmaß", "Passepartout", "Preis EUR", "Rahmen im Preis",
    "Versicherungswert EUR", "Status", "Käufer:in", "Notiz",
]


def zahl(v):
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return int(v) if float(v).is_integer() else float(v)
    try:
        return float(str(v).replace(",", "."))
    except ValueError:
        return None


def fmt_zahl(v):
    n = zahl(v)
    if n is None:
        return ""
    return str(int(n)) if float(n).is_integer() else f"{n:.1f}".replace(".", ",")


_GRAMMATUR = re.compile(
    r"\s*,?\s*\d+(?:[.,]\d+)?(?:\s*[–-]\s*\d+(?:[.,]\d+)?)?\s*g(?:\s*/?\s*m[²2])?\b")


def _ohne_grammatur(papier):
    """Entfernt die Flächengewicht-Angabe aus dem Papiernamen — „325 g/m²“,
    „300g/m²“, „115–120 g/m²“, aber auch die Kurzform „308 g“. Auf den
    Schildern und in den Listen nicht nötig und sorgt für unnötige Umbrüche."""
    t = _GRAMMATUR.sub("", papier)
    t = re.sub(r"\s{2,}", " ", t).strip(" ,")
    return re.sub(r"\(\s*\)", "", t).strip(" ,")


def lade_werke(pfad):
    if not os.path.exists(pfad):
        sys.exit(f"FEHLER: {pfad} nicht gefunden.")
    wb = load_workbook(pfad, data_only=True)
    if "Werke" not in wb.sheetnames:
        sys.exit("FEHLER: Blatt 'Werke' fehlt in der Datei.")
    ws = wb["Werke"]
    header = [c.value for c in ws[1]]
    idx = {h: i for i, h in enumerate(header) if h}

    fehlend = [s for s in SPALTEN if s not in idx]
    if fehlend:
        sys.exit("FEHLER: Spalten fehlen: " + ", ".join(fehlend))

    werke = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        titel = row[idx["Titel"]]
        kuenstler = row[idx["Künstler:in"]]
        if not titel and not kuenstler:
            continue
        w = {}
        for s in SPALTEN:
            v = row[idx[s]]
            w[s] = "" if v is None else v
        # „40x50" -> „40×50" in den Maß-Textfeldern (falls von Hand getippt)
        for feld in ("Rahmenmaß", "Rahmenmodell"):
            if w.get(feld):
                w[feld] = mal_zeichen(w[feld])
        # abgeleitete Felder — unabhängig davon, ob Excel neu gerechnet hat
        h, b = fmt_zahl(w["Bildmaß H cm"]), fmt_zahl(w["Bildmaß B cm"])
        w["_masse"] = f"{h} × {b} cm" if h and b else str(w["Maßangabe"] or "")
        bh, bb = fmt_zahl(w["Blattmaß H cm"]), fmt_zahl(w["Blattmaß B cm"])
        w["_blattmass"] = f"{bh} × {bb} cm" if bh and bb else ""
        au, ap = zahl(w["Auflage"]), zahl(w["AP"])
        if au:
            w["_auflage"] = f"Auflage {int(au)}" + (f" (+{int(ap)} AP)" if ap else "")
        else:
            w["_auflage"] = str(w["Auflageangabe"] or "")
        ort = str(w["Ort"] or "").strip()
        land = str(w["Land"] or "").strip()
        jahr = zahl(w["Jahr"])
        w["_ortjahr"] = ", ".join([x for x in [ort, land, (str(int(jahr)) if jahr else "")] if x])
        verf = str(w["Druckverfahren"] or "").strip()
        pap = _ohne_grammatur(str(w["Papier"] or "").strip())
        w["_technik"] = ", ".join([x for x in [verf, pap] if x])
        w["_nr"] = str(w["Nr"] or "").strip()
        w["_status"] = str(w["Status"] or "").strip().lower()
        w["_preis"] = zahl(w["Preis EUR"])
        w["_rahmen_inkl"] = str(w["Rahmen im Preis"] or "").strip().lower()
        w["_vers"] = zahl(w["Versicherungswert EUR"])
        # optionale Spalte "Beispiel": markiert die 14 Demowerke aus build_master.py
        bsp = row[idx["Beispiel"]] if "Beispiel" in idx else None
        w["_beispiel"] = str(bsp or "").strip().lower() in ("ja", "yes", "true", "1", "x")
        werke.append(w)
    return werke


def pruefe(werke):
    """Meldet fehlende Pflichtfelder, bevor gedruckt wird."""
    pflicht_schild = ["Nr", "Künstler:in", "Titel", "Ort", "Land", "Jahr",
                      "Druckverfahren", "Papier", "Bildmaß H cm", "Bildmaß B cm", "Auflage"]
    probleme = []
    nummern = {}
    for w in werke:
        fehlt = [f for f in pflicht_schild if w.get(f) in ("", None)]
        if fehlt:
            probleme.append(f"  Nr {w['_nr'] or '???'} „{w['Titel']}“: fehlt {', '.join(fehlt)}")
        if w["_status"] not in ("nicht verkäuflich", "") and w["_preis"] is None:
            probleme.append(f"  Nr {w['_nr']} „{w['Titel']}“: kein Preis, Status ist „{w['Status']}“")
        if w["_preis"] is not None and w["_rahmen_inkl"] not in ("ja", "nein"):
            probleme.append(f"  Nr {w['_nr']} „{w['Titel']}“: „Rahmen im Preis“ ist weder ja noch nein")
        nummern.setdefault(w["_nr"], []).append(w["Titel"])
    for nr, titel in nummern.items():
        if len(titel) > 1:
            probleme.append(f"  Werknummer {nr} doppelt vergeben: {', '.join(str(t) for t in titel)}")
    return probleme


# ----------------------------------------------------------------- Textwerkzeug
def wrap(text, font, size, maxw):
    """Bricht Text auf maxw um und gibt eine Liste von Zeilen zurück."""
    text = str(text or "").strip()
    if not text:
        return []
    words, lines, cur = text.split(), [], ""
    for wd in words:
        probe = (cur + " " + wd).strip()
        if pdfmetrics.stringWidth(probe, font, size) <= maxw or not cur:
            cur = probe
        else:
            lines.append(cur)
            cur = wd
    if cur:
        lines.append(cur)
    return lines


def draw_wrapped(c, x, y, text, font, size, maxw, leading, maxlines=None, color=(0, 0, 0)):
    lines = wrap(text, font, size, maxw)
    if maxlines and len(lines) > maxlines:
        lines = lines[:maxlines]
        while lines and pdfmetrics.stringWidth(lines[-1] + " …", font, size) > maxw:
            lines[-1] = lines[-1].rsplit(" ", 1)[0]
        lines[-1] += " …"
    c.setFont(font, size)
    c.setFillColorRGB(*color)
    for ln in lines:
        c.drawString(x, y, ln)
        y -= leading
    return y


# ----------------------------------------------------------------- 01 Wandschilder
# DIN-A-Reihe in mm (Hochformat: Breite, Höhe). Wandschilder stehen quer, die
# lange Seite ist also die Breite.
DIN_MM = {0: (841, 1189), 1: (594, 841), 2: (420, 594), 3: (297, 420), 4: (210, 297),
          5: (148, 210), 6: (105, 148), 7: (74, 105), 8: (52, 74), 9: (37, 52), 10: (26, 37)}
FORMATE = [f"a{n}" for n in DIN_MM]                    # a0 … a10

# Bezugsgröße der automatischen Skalierung: A8 quer. Das ist die kompakteste
# Variante, bei der alle Angaben noch ohne Kürzung passen (QR mit 9 mm knapp,
# aber mit guter Druckauflösung scannbar). Alle anderen DIN-Formate, für die
# unten kein eigener Eintrag steht, werden linear davon abgeleitet.
REF = dict(w=74 * mm, h=52 * mm, pad=4.5 * mm, qr=9 * mm,
           s_nr=8.5, s_art=8.5, s_tit=10, s_ort=7.5, s_tec=6.6, s_hint=6.2)
_SKALIERT = ("pad", "qr", "s_nr", "s_art", "s_tit", "s_ort", "s_tec", "s_hint")

# Feste, von Hand abgestimmte Bogenformate. A7 und A6 sind bewusst etwas
# kleiner als das DIN-Maß, damit 8 bzw. 4 Schilder samt Schnittmarken auf einen
# A4-Bogen passen. Alle anderen Formate (a0–a5, a9, a10) sind exakt DIN.
PRESETS = {
    # sehr klein: 10 Schilder auf einem A4-Bogen hoch
    "a8": dict(REF, cols=2, rows=5, seite=A4),
    # klein: 8 Schilder auf einem A4-Bogen hoch
    "a7": dict(w=100 * mm, h=70 * mm, cols=2, rows=4, seite=A4, pad=7.0 * mm,
               s_nr=10.5, s_art=11, s_tit=13, s_ort=9.5, s_tec=8.3, s_hint=7.6,
               qr=12 * mm),
    # groß und barrierearm: 4 Schilder auf einem A4-Bogen QUER
    "a6": dict(w=140 * mm, h=100 * mm, cols=2, rows=2, seite=landscape(A4), pad=12 * mm,
               s_nr=13, s_art=15, s_tit=18, s_ort=13, s_tec=11, s_hint=10,
               qr=17 * mm),
}

_SCHRIFT_KEYS = (("nummer", "s_nr"), ("kuenstlerin", "s_art"), ("titel", "s_tit"),
                 ("ortjahr", "s_ort"), ("technik", "s_tec"), ("hinweis", "s_hint"))


def _din_seite(name):
    """'A4' / 'a4quer' / 'A3hoch' -> ((Breite, Höhe) in pt, 'quer'|'hoch'|None); None bei Unsinn.
    'A4' ohne Zusatz gilt als Hochformat (wie bisher in design.json)."""
    m = re.fullmatch(r"a(\d+)(quer|hoch)?", str(name).strip().lower())
    if not m or int(m.group(1)) not in DIN_MM:
        return None
    b, h = DIN_MM[int(m.group(1))]
    return ((h * mm, b * mm) if m.group(2) == "quer" else (b * mm, h * mm)), m.group(2)


def seitenname(seite):
    """(Breite, Höhe) in pt -> 'A4 hoch' / 'A3 quer' / '123 × 456 mm'."""
    b, h = seite[0] / mm, seite[1] / mm
    for n, (db, dh) in DIN_MM.items():
        if abs(b - db) < 1 and abs(h - dh) < 1:
            return f"A{n} hoch"
        if abs(b - dh) < 1 and abs(h - db) < 1:
            return f"A{n} quer"
    return f"{b:.0f} × {h:.0f} mm"


def _skaliere(cfg, w, h):
    """Schild auf die Größe w × h bringen: Schriften, Rand und QR wachsen bzw.
    schrumpfen im selben Verhältnis (kleinere Seite des Verhältnisses gewinnt)."""
    k = min(w / cfg["w"], h / cfg["h"])
    neu = dict(cfg, w=w, h=h)
    for key in _SKALIERT:
        neu[key] = cfg[key] * k
    return neu


def plane_bogen(w, h, rand, bogen=None):
    """Wie viele Schilder w × h passen (mit Rand für Schnittmarken) auf den Bogen?
    bogen=None/'auto': A4, hoch oder quer — was mehr Schilder ergibt.
    bogen='A3', 'A4quer' …: genau dieser Bogen (ohne Zusatz = hoch).
    -> (seite, spalten, zeilen) oder None, wenn nicht einmal eines passt
       (dann gilt: Blatt = Karte, siehe schild_cfg)."""
    if bogen and str(bogen).lower() != "auto":
        p = _din_seite(bogen)
        if not p:
            sys.exit(f"FEHLER: design.json → bogen „{bogen}“ ist kein DIN-A-Bogen "
                     f"(z. B. A4, A4quer, A3).")
        seiten = [p[0]]
    else:
        seiten = [A4, landscape(A4)]
    bester = None
    for sw, sh in seiten:
        cols = int((sw - 2 * rand + 0.01 * mm) // w)
        rows = int((sh - 2 * rand + 0.01 * mm) // h)
        if cols >= 1 and rows >= 1 and (bester is None or cols * rows > bester[1] * bester[2]):
            bester = ((sw, sh), cols, rows)
    return bester


def schild_cfg(fmt):
    """Endgültige Konfiguration eines Schildformats (a0 … a10):
    Grundwerte (eigener Eintrag in PRESETS oder linear von REF abgeleitet)
    → Überschreibungen aus design.json → Anordnung auf dem Bogen."""
    fmt = fmt.lower()
    if fmt not in FORMATE:
        sys.exit(f"FEHLER: unbekanntes Schildformat „{fmt}“ (erlaubt: {', '.join(FORMATE)}).")
    ov = DESIGN_SCHILD.get(fmt, {})
    rand = ANPASSUNG["bogenrand_mm"] * mm

    if fmt in PRESETS:
        cfg = dict(PRESETS[fmt])
    else:
        b, h = DIN_MM[int(fmt[1:])]
        cfg = _skaliere(dict(REF), h * mm, b * mm)          # quer: lange Seite = Breite

    if "breite_mm" in ov or "hoehe_mm" in ov:
        cfg = _skaliere(cfg, ov["breite_mm"] * mm if "breite_mm" in ov else cfg["w"],
                        ov["hoehe_mm"] * mm if "hoehe_mm" in ov else cfg["h"])
        for k in ("cols", "rows", "seite"):                  # Raster passt nicht mehr → neu rechnen
            cfg.pop(k, None)
    if "rand_mm" in ov:
        cfg["pad"] = ov["rand_mm"] * mm
    if "qr_mm" in ov:
        cfg["qr"] = ov["qr_mm"] * mm
    for k, ziel in _SCHRIFT_KEYS:
        if k in ov.get("schriftgroessen_pt", {}):
            cfg[ziel] = ov["schriftgroessen_pt"][k]

    # --- Anordnung auf dem Bogen ---------------------------------------------
    cfg["blatt"] = False
    if str(ov.get("bogen") or "").lower() in ("karte", "blatt"):
        # ausdrücklich: eine Karte pro Blatt in Schildgröße, keine Schnittmarken
        cfg.update(blatt=True, seite=(cfg["w"], cfg["h"]), cols=1, rows=1)
    elif ov.get("bogen") or "seite" not in cfg:
        plan = plane_bogen(cfg["w"], cfg["h"], rand, ov.get("bogen"))
        if plan is None:
            # Schild passt (mit Rand) auf keinen A4-Bogen: Blatt = Karte, eine
            # Karte pro Seite in Schildgröße, keine Schnittmarken.
            cfg.update(blatt=True, seite=(cfg["w"], cfg["h"]), cols=1, rows=1)
        else:
            cfg["seite"] = plan[0]
            cfg["cols"], cfg["rows"] = plan[1], plan[2]
    if not cfg["blatt"]:
        if "spalten" in ov:
            cfg["cols"] = int(ov["spalten"])
        if "zeilen" in ov:
            cfg["rows"] = int(ov["zeilen"])
    cfg["fmt"] = fmt
    return cfg


# ------------------------------------------------ Inhalt an das Schild anpassen
# Das Raster in schild_zeichnen() ist fest (Nr, Name, 2× Titel, Ort, Strich,
# 2× Technik, Maß/Auflage, Hinweis): seine Höhe hängt nur von den Schrift-
# größen ab, nicht vom Text. Passt sie nicht in das Schild, wird alles im
# selben Verhältnis verkleinert (fit_faktor). Zu lange Texte schrumpfen einzeln
# bis text_min_faktor und werden erst danach mit „…“ gekürzt.
def _raster_hoehe(cfg, f=1.0):
    """Höhe in pt, die das Raster von schild_zeichnen() bis zur Unterkante der
    letzten Zeile belegt, bei Skalierung f."""
    g = cfg["s_art"] * f / 11.0
    zeilen = (cfg["s_nr"] * 1.2 + cfg["s_art"] * 1.25 + 2 * cfg["s_tit"] * 1.22
              + cfg["s_ort"] * 1.3 + 3 * cfg["s_tec"] * 1.34 + cfg["s_hint"] * 1.2
              - cfg["s_hint"] * 0.1) * f
    luecken = (3.0 + 1.0 + 1.0 + 2.2 + 3.2 + 3.5) * mm * g
    return zeilen + luecken


def fit_faktor(cfg, bh):
    """Skalierung (≤ 1), mit der das Raster in die Schildhöhe bh passt."""
    if not ANPASSUNG["aktiv"]:
        return 1.0
    frei = bh - 2 * cfg["pad"]
    f = min(1.0, frei / _raster_hoehe(cfg))
    return max(f, ANPASSUNG["inhalt_min_faktor"])


def _mit_punkten(text, font, size, maxw):
    t = text
    while t and pdfmetrics.stringWidth(t + "…", font, size) > maxw:
        t = t[:-1]
    return t.rstrip() + "…"


def einzeilig(text, font, size, maxw, min_faktor=0.7, schritt=0.25):
    """Text in EINE Zeile bringen: erst Schrift bis min_faktor verkleinern,
    dann mit „…“ kürzen. -> (Zeile, Schriftgröße)"""
    text = str(text or "").strip()
    if not text:
        return "", size
    s = size
    while pdfmetrics.stringWidth(text, font, s) > maxw and s - schritt >= size * min_faktor:
        s -= schritt
    if pdfmetrics.stringWidth(text, font, s) > maxw:
        text = _mit_punkten(text, font, s, maxw)
    return text, s


def mehrzeilig(text, font, size, maxw, maxzeilen, min_faktor=0.7, schritt=0.25):
    """Text auf höchstens maxzeilen umbrechen: erst Schrift bis min_faktor
    verkleinern, dann mit „…“ kürzen. -> ([Zeilen], Schriftgröße)"""
    text = str(text or "").strip()
    if not text:
        return [], size
    s = size
    zeilen = wrap(text, font, s, maxw)
    while len(zeilen) > maxzeilen and s - schritt >= size * min_faktor:
        s -= schritt
        zeilen = wrap(text, font, s, maxw)
    if len(zeilen) > maxzeilen:
        zeilen = zeilen[:maxzeilen]
        zeilen[-1] = _mit_punkten(zeilen[-1] + " …", font, s, maxw)
    zeilen = [z if pdfmetrics.stringWidth(z, font, s) <= maxw else _mit_punkten(z, font, s, maxw)
              for z in zeilen]
    return zeilen, s


def schild_zeichnen(c, x, y, bw, bh, w, cfg, preishinweis=True, statuspunkt=False,
                    qr=None):
    """Zeichnet ein Wandschild in das Rechteck mit unterer linker Ecke (x, y).

    Festes Raster: jede Angabe steht immer auf derselben Zeile, unabhängig von
    der Textlänge der anderen Felder. Fehlt ein Wert, bleibt seine Zeile leer —
    so sehen alle Schilder gleich aus. Titel und Technik/Papier haben je zwei
    reservierte Zeilen. QR-Code (falls vorhanden) unten rechts.

    Größen passen sich an: passt das Raster nicht in die Schildhöhe, wird
    alles verkleinert (fit_faktor); zu lange Texte schrumpfen einzeln oder
    werden gekürzt (einzeilig / mehrzeilig).
    """
    pad = cfg["pad"]
    tw = bw - 2 * pad
    qs = cfg.get("qr", 12 * mm) if qr else 0
    tw_lo = tw - (qs + 3 * mm) if qr else tw     # schmalere Breite für die unteren Zeilen
    grau = FARBEN["sekundaer"]

    f = fit_faktor(cfg, bh)
    s_nr, s_art, s_tit0, s_ort, s_tec, s_hint = (cfg[k] * f for k in
                                                  ("s_nr", "s_art", "s_tit", "s_ort", "s_tec", "s_hint"))
    g = s_art / 11.0                              # skaliert die mm-Abstände je Format
    tmin = ANPASSUNG["text_min_faktor"] if ANPASSUNG["aktiv"] else 1.0

    # --- messen -----------------------------------------------------------
    art, s_art_z = einzeilig(w["Künstler:in"], F_BOLD, s_art, tw, tmin)
    l_tit, s_tit = mehrzeilig(w["Titel"], F_IT, s_tit0, tw, 2, tmin)
    ort, s_ort_z = einzeilig(w["_ortjahr"], F_REG, s_ort, tw, tmin)
    l_tec, s_tec_z = mehrzeilig(w["_technik"], F_REG, s_tec, tw_lo, 2, tmin)
    det = " · ".join([z for z in [w["_masse"], w["_auflage"]] if z])
    det, s_det_z = einzeilig(det, F_REG, s_tec, tw_lo, tmin)

    hint = ""
    if w["_status"] == "nicht verkäuflich":
        hint = "unverkäuflich"
    elif preishinweis:
        hint = "Preis siehe Werkliste"
    hint, s_hint_z = einzeilig(hint, F_REG, s_hint, tw_lo, tmin)

    # --- zeichnen: festes Raster von oben ---------------------------------
    cy = y + bh - pad

    def block(lines, font, size, farbe, lead, reserve, gap_after=0.0):
        nonlocal cy
        c.setFont(font, size)
        c.setFillColorRGB(*farbe)
        for i in range(reserve):
            cy -= size * 0.92
            if i < len(lines) and lines[i]:
                c.drawString(x + pad, cy, lines[i])
            cy -= lead - size * 0.92
        cy -= gap_after

    block([w["_nr"]], F_BOLD, s_nr, grau, s_nr * 1.2, 1, 3.0 * mm * g)
    block([art], F_BOLD, s_art_z, FARBEN["text"], s_art * 1.25, 1, 1.0 * mm * g)
    block(l_tit, F_IT, s_tit, FARBEN["text"], s_tit0 * 1.22, 2, 1.0 * mm * g)
    block([ort], F_REG, s_ort_z, FARBEN["text"], s_ort * 1.3, 1, 2.2 * mm * g)

    c.setStrokeColorRGB(*FARBEN["linie"])
    c.setLineWidth(max(0.25, 0.5 * f))
    c.line(x + pad, cy, x + pad + tw * 0.28, cy)
    cy -= 3.2 * mm * g

    block(l_tec, F_REG, s_tec_z, grau, s_tec * 1.34, 2)
    block([det], F_REG, s_det_z, grau, s_tec * 1.34, 1, 3.5 * mm * g)
    block([hint], F_REG, s_hint_z, grau, s_hint * 1.2, 1)

    if qr:
        qr_zeichnen(c, qr, x + bw - pad - qs, y + pad, qs)

    if statuspunkt and w["_status"] in ("verkauft", "reserviert"):
        r = 1.7 * mm * max(g, 1.0)
        c.setFillColorRGB(*(FARBEN["signal"] if w["_status"] == "verkauft"
                            else (0.95, 0.72, 0.10)))
        py = y + pad + (qs + 2.5 * mm * max(g, 1.0) + r if qr else 0)
        c.circle(x + bw - pad - r, py, r, stroke=0, fill=1)

    c.setFillColorRGB(0, 0, 0)


def schild_hinweise(cfg):
    """Zeilen für die Konsole: Größe, Anordnung, Anpassungen."""
    bw, bh = cfg["w"], cfg["h"]
    if cfg["blatt"]:
        zeilen = [f"{bw / mm:.0f} × {bh / mm:.0f} mm, eine Karte pro Blatt (Blatt = Karte, "
                  f"keine Schnittmarken)"]
    else:
        zeilen = [f"{bw / mm:.0f} × {bh / mm:.0f} mm, {cfg['cols']} × {cfg['rows']} = "
                  f"{cfg['cols'] * cfg['rows']} je Bogen {seitenname(cfg['seite'])}"]
        pw, ph = cfg["seite"]
        if cfg["cols"] * bw > pw + 0.01 * mm or cfg["rows"] * bh > ph + 0.01 * mm:
            zeilen.append("ACHTUNG: Raster ist größer als der Bogen — spalten/zeilen/bogen "
                          "in design.json prüfen")
    f = fit_faktor(cfg, bh)
    if f < 0.995:
        zeilen.append(f"Inhalt auf {f * 100:.0f} % verkleinert, damit das Raster ins Schild passt")
        if f <= ANPASSUNG["inhalt_min_faktor"] + 1e-9:
            zeilen.append("ACHTUNG: Untergrenze inhalt_min_faktor erreicht — Inhalt kann überlaufen")
    kleinste = cfg["s_hint"] * f
    if kleinste < 5:
        zeilen.append(f"ACHTUNG: kleinste Schrift nur {kleinste:.1f} pt — nur bei hoher "
                      f"Druckauflösung lesbar")
    return zeilen


def pdf_wandschilder(werke, pfad, format_="a8", preishinweis=True, statuspunkt=False,
                     marken=True, insta=None):
    cfg = schild_cfg(format_)
    insta = insta or {}
    bw, bh, cols, rows = cfg["w"], cfg["h"], cfg["cols"], cfg["rows"]
    pw, ph = cfg["seite"]
    if cfg["blatt"]:
        marken = False
    mx = (pw - cols * bw) / 2
    my = (ph - rows * bh) / 2

    c = canvas.Canvas(pfad, pagesize=A4)
    c.setTitle("FOTOTAGE MUSTERSTADT 2026 — Wandschilder")
    per = cols * rows
    zeichne_deckblatt(
        c, "Wandschilder", "Die Beschriftung an der Wand",
        "Neben jeder Arbeit hängt ein Schild mit Künstler:in, Titel, Ort und Jahr, "
        "Technik, Maß und Auflage. Der QR-Code führt auf das Instagram-Profil der "
        "Künstlerin oder des Künstlers.",
    )
    c.showPage()
    c.setPageSize(cfg["seite"])
    for i, w in enumerate(werke):
        if i % per == 0:
            if i:
                c.showPage()
            if marken:
                gx0, gx1 = mx, mx + cols * bw
                gy0, gy1 = my, my + rows * bh
                # durchgehende gestrichelte Schnittlinien über den ganzen Bogen
                c.setStrokeColorRGB(0.62, 0.62, 0.62)
                c.setLineWidth(0.35)
                c.setDash(2.2, 2.2)
                for cc in range(cols + 1):
                    xx = mx + cc * bw
                    c.line(xx, gy0 - 2 * mm, xx, gy1 + 2 * mm)
                for rr in range(rows + 1):
                    yy = my + rr * bh
                    c.line(gx0 - 2 * mm, yy, gx1 + 2 * mm, yy)
                c.setDash()
                # kurze solide Passermarken außen für den exakten Anschnitt
                c.setStrokeColorRGB(0.55, 0.55, 0.55)
                c.setLineWidth(0.3)
                for cc in range(cols + 1):
                    xx = mx + cc * bw
                    c.line(xx, gy0 - 5.5 * mm, xx, gy0 - 2.5 * mm)
                    c.line(xx, gy1 + 2.5 * mm, xx, gy1 + 5.5 * mm)
                for rr in range(rows + 1):
                    yy = my + rr * bh
                    c.line(gx0 - 5.5 * mm, yy, gx0 - 2.5 * mm, yy)
                    c.line(gx1 + 2.5 * mm, yy, gx1 + 5.5 * mm, yy)
                c.setFont(F_REG, 5.5)
                c.setFillColorRGB(0.72, 0.72, 0.72)
                c.drawString(mx, 3 * mm,
                             f"FOTOTAGE MUSTERSTADT 2026 · Wandschilder {format_.upper()} · "
                             f"an den gestrichelten Linien schneiden · Bogen {i // per + 1}")
        k = i % per
        col, row = k % cols, k // cols
        x = mx + col * bw
        y = my + (rows - 1 - row) * bh
        png = qr_png(insta.get(str(w["Künstler:in"]).strip(), ""))
        schild_zeichnen(c, x, y, bw, bh, w, cfg, preishinweis, statuspunkt, qr=png)
    c.save()
    return len(werke), (len(werke) + per - 1) // per, schild_hinweise(cfg)


# --------------------------------------------- Einzelkarte (ein Schild pro Blatt)
# Für sehr große Wandbilder, bei denen ein kleines Schild optisch untergeht.
# Eine Karte pro Blatt in Kartengröße (z. B. A3 quer), kein Deckblatt, keine
# Schnittmarken — das Blatt selbst ist die Karte.
def pdf_wandschild_einzel(werk, pfad, format_="a3", preishinweis=True, statuspunkt=False,
                          insta=None):
    cfg = schild_cfg(format_)
    seite = (cfg["w"], cfg["h"])
    c = canvas.Canvas(pfad, pagesize=seite)
    c.setTitle(f"FOTOTAGE MUSTERSTADT 2026 — Wandschild Nr {werk['_nr']} ({format_.upper()})")
    png = qr_png((insta or {}).get(str(werk["Künstler:in"]).strip(), ""))
    schild_zeichnen(c, 0, 0, seite[0], seite[1], werk, cfg,
                    preishinweis, statuspunkt, qr=png)
    c.save()
    return schild_hinweise(dict(cfg, blatt=True))

# ----------------------------------------------------------------- 02 Werkliste
def preis_text(w):
    if w["_status"] == "verkauft":
        return "verkauft"
    if w["_status"] == "nicht verkäuflich":
        return "unverkäuflich"
    if w["_preis"] is None:
        return "auf Anfrage"
    p = f"{int(w['_preis']):,}".replace(",", ".") + " €"
    if w["_rahmen_inkl"] == "ja":
        p += " inkl. Rahmen"
    elif w["_rahmen_inkl"] == "nein":
        p += " zzgl. Rahmen"
    if w["_status"] == "reserviert":
        p += " · reserviert"
    return p


def pdf_werkliste(werke, pfad, spalten=2, gruppe="Wand"):
    pw, ph = A4
    ml, mr, mt, mb = 16 * mm, 16 * mm, 20 * mm, 20 * mm
    gutter = 8 * mm
    colw = (pw - ml - mr - (spalten - 1) * gutter) / spalten

    TXT, SEK, TER = FARBEN["text"], FARBEN["sekundaer"], FARBEN["tertiaer"]
    LIN, SIG = FARBEN["linie"], FARBEN["signal"]

    c = canvas.Canvas(pfad, pagesize=A4)
    c.setTitle("FOTOTAGE MUSTERSTADT 2026 — Werkliste")

    zeichne_deckblatt(
        c, "Werkliste", "Alle Arbeiten mit Preisen",
        "Alle ausgestellten Arbeiten in der Reihenfolge der Wandschilder, mit Preis "
        "und Angaben zu Auflage und Rahmung. Verkäufe und Reservierungen sind "
        "gekennzeichnet. Kaufinteresse an der Theke, direkt bei der Künstlerin oder "
        "dem Künstler oder unter kontakt@beispiel-verein.de.",
    )
    c.showPage()

    state = {"page": 0, "col": 0, "y": 0.0}

    def kopf(erste):
        c.setFillColorRGB(*TXT)
        if erste:
            c.setFont(F_BOLD, 20)
            c.drawString(ml, ph - mt + 4 * mm, AUSSTELLUNG["titel"])
            c.setFont(F_REG, 10)
            c.setFillColorRGB(*SEK)
            c.drawString(ml, ph - mt - 2.5 * mm, "Werkliste · " + AUSSTELLUNG["untertitel"])
            c.drawString(ml, ph - mt - 7 * mm,
                         AUSSTELLUNG["laufzeit"] + " · " + AUSSTELLUNG["ort"])
            c.setStrokeColorRGB(*TXT)
            c.setLineWidth(0.8)
            c.line(ml, ph - mt - 11 * mm, pw - mr, ph - mt - 11 * mm)
            return ph - mt - 20 * mm
        c.setFont(F_REG, 8)
        c.setFillColorRGB(*TER)
        c.drawString(ml, ph - mt + 2 * mm, AUSSTELLUNG["titel"] + " · Werkliste")
        c.setStrokeColorRGB(*LIN)
        c.setLineWidth(0.4)
        c.line(ml, ph - mt - 1.5 * mm, pw - mr, ph - mt - 1.5 * mm)
        return ph - mt - 8 * mm

    def fuss():
        c.setFont(F_REG, 7)
        c.setFillColorRGB(*TER)
        c.drawString(ml, mb - 8 * mm,
                     f"{AUSSTELLUNG['web']} · {AUSSTELLUNG['mail']}")
        c.drawRightString(pw - mr, mb - 8 * mm, f"Seite {state['page']}")

    def neue_seite():
        if state["page"]:
            fuss()
            c.showPage()
        state["page"] += 1
        state["col"] = 0
        state["y"] = kopf(state["page"] == 1)

    def spalte_x():
        return ml + state["col"] * (colw + gutter)

    def platz(hoehe):
        if state["y"] - hoehe >= mb:
            return
        if state["col"] < spalten - 1:
            state["col"] += 1
            state["y"] = ph - mt - (20 * mm if state["page"] == 1 else 8 * mm)
        else:
            neue_seite()

    NAME_ABSTAND = 1.6 * mm          # Luft zwischen Ausstellende:r und Titel
    KARTE_PAD = 3.2 * mm             # Innenabstand der Karte (jede Zeile = eigene Karte)
    KARTE_GAP = 3.0 * mm             # Abstand zwischen zwei Karten

    def eintrag_hoehe(w):
        tw = colw - 2 * KARTE_PAD - 11 * mm
        h = 0
        h += 11.5 * 1.15 + NAME_ABSTAND                      # Künstler:in
        h += len(wrap(w["Titel"], F_IT, 10.5, tw)) * 10.5 * 1.2
        h += len(wrap(w["_ortjahr"] + " · " + w["_technik"], F_REG, 8, tw)) * 8 * 1.28
        det = " · ".join([z for z in [w["_masse"], w["_auflage"]] if z])
        h += len(wrap(det, F_REG, 8, tw)) * 8 * 1.28
        h += 9 * 1.3
        return h + 2 * KARTE_PAD + KARTE_GAP

    def eintrag(w, gname=None):
        h = eintrag_hoehe(w)
        if state["y"] - h < mb and gname:
            platz(h)
            zeichne_gruppentitel(f"{gname} (Fortsetzung)")
        else:
            platz(h)
        x = spalte_x()
        top = state["y"]
        boden = top - (h - KARTE_GAP)

        c.setStrokeColorRGB(*LIN)
        c.setLineWidth(0.4)
        c.roundRect(x, boden, colw, top - boden, 1.6 * mm, stroke=1, fill=0)

        x0 = x + KARTE_PAD
        xi = x0 + 11 * mm
        tw = colw - 2 * KARTE_PAD - 11 * mm
        y = top - KARTE_PAD

        c.setFont(F_BOLD, 9)
        c.setFillColorRGB(*TER)
        c.drawString(x0, y - 9, w["_nr"])

        c.setFont(F_BOLD, 11.5)
        c.setFillColorRGB(*TXT)
        c.drawString(xi, y - 11.5 * 0.85, w["Künstler:in"])
        y -= 11.5 * 1.15 + NAME_ABSTAND

        y = draw_wrapped(c, xi, y - 10.5 * 0.6, w["Titel"], F_IT, 10.5, tw, 10.5 * 1.2,
                         color=TXT)
        y -= 0.6 * mm
        y = draw_wrapped(c, xi, y, w["_ortjahr"] + " · " + w["_technik"], F_REG, 8, tw,
                         8 * 1.28, color=SEK)
        det = " · ".join([z for z in [w["_masse"], w["_auflage"]] if z])
        if w["_blattmass"]:
            det = det.replace(w["_masse"], f"{w['_masse']} (Blatt {w['_blattmass']})", 1)
        y = draw_wrapped(c, xi, y, det, F_REG, 8, tw, 8 * 1.28, color=SEK)

        y -= 0.8 * mm
        c.setFont(F_BOLD, 9)
        pt = preis_text(w)
        c.setFillColorRGB(*(SIG if w["_status"] in ("verkauft", "reserviert") else TXT))
        c.drawString(xi, y - 9 * 0.85, pt)

        state["y"] = boden - KARTE_GAP
        c.setFillColorRGB(*TXT)

    def gruppentitel(t):
        platz(20 * mm)
        zeichne_gruppentitel(t)

    def zeichne_gruppentitel(t):
        x = spalte_x()
        y = state["y"]
        c.setFont(F_BOLD, 10)
        c.setFillColorRGB(*TXT)
        c.drawString(x, y - 10, t)
        c.setStrokeColorRGB(*TXT)
        c.setLineWidth(0.6)
        c.line(x, y - 13.5, x + colw, y - 13.5)
        state["y"] = y - 19

    neue_seite()

    ohne_gruppen = gruppe in (None, "", "ohne", "keine")

    def sortkey(w):
        if ohne_gruppen:
            return (w["_nr"] or "zzz", w["Künstler:in"])
        g = str(w.get(gruppe, "") or "")
        p = zahl(w.get("Position")) or 0
        return (g, p, w["_nr"])

    aktuell = None
    for w in sorted(werke, key=sortkey):
        if ohne_gruppen:
            eintrag(w)
            continue
        g = str(w.get(gruppe, "") or "")
        name = f"Wand {g}" if gruppe == "Wand" and g else (g or "Ohne Zuordnung")
        if g != aktuell:
            aktuell = g
            gruppentitel(name)
        eintrag(w, name)

    fuss()
    c.save()
    return state["page"]


# ----------------------------------------------------------------- 03 Rückseitenetiketten
def pdf_rueckseiten(werke, pfad, teilnehmende=None):
    teilnehmende = teilnehmende or {}
    pw, ph = A4
    bw, bh = 90 * mm, 45 * mm
    cols, rows = 2, 6
    mx = (pw - cols * bw) / 2
    my = (ph - rows * bh) / 2
    per = cols * rows

    c = canvas.Canvas(pfad, pagesize=A4)
    c.setTitle("FOTOTAGE MUSTERSTADT 2026 — Rückseitenetiketten")
    zeichne_deckblatt(
        c, "Rückseitenetiketten", "Für die Rückseite der Rahmen",
        "Je ein Etikett pro Arbeit mit Werknummer, Künstler:in, Titel, Maß, Technik, "
        "Rahmenmodell, Versicherungswert und Kontakt — für die Rückseite von Rahmen "
        "oder Passepartout.",
    )
    c.showPage()
    for i, w in enumerate(werke):
        if i % per == 0:
            if i:
                c.showPage()
            c.setStrokeColorRGB(0.80, 0.80, 0.80)
            c.setLineWidth(0.25)
            for cc in range(cols + 1):
                xx = mx + cc * bw
                c.line(xx, my - 4 * mm, xx, my - 1 * mm)
                c.line(xx, my + rows * bh + 1 * mm, xx, my + rows * bh + 4 * mm)
            for rr in range(rows + 1):
                yy = my + rr * bh
                c.line(mx - 4 * mm, yy, mx - 1 * mm, yy)
                c.line(mx + cols * bw + 1 * mm, yy, mx + cols * bw + 4 * mm, yy)
        k = i % per
        x = mx + (k % cols) * bw
        y = my + (rows - 1 - k // cols) * bh
        pad = 5 * mm
        tw = bw - 2 * pad
        cy = y + bh - pad

        c.setStrokeColorRGB(0.85, 0.85, 0.85)
        c.setLineWidth(0.3)
        c.rect(x + 2 * mm, y + 2 * mm, bw - 4 * mm, bh - 4 * mm, stroke=1, fill=0)

        c.setFont(F_REG, 6)
        c.setFillColorRGB(0.55, 0.55, 0.55)
        c.drawString(x + pad, cy - 6, AUSSTELLUNG["titel"] + " · Werk Nr. " + w["_nr"])
        cy -= 6 + 2.2 * mm

        cy -= 9 * 0.85
        c.setFont(F_BOLD, 9)
        c.setFillColorRGB(0, 0, 0)
        c.drawString(x + pad, cy, str(w["Künstler:in"]))
        cy -= 2.8 * mm

        cy = draw_wrapped(c, x + pad, cy, w["Titel"], F_IT, 8.5, tw, 8.5 * 1.2, maxlines=1)
        cy -= 0.6 * mm
        cy = draw_wrapped(c, x + pad, cy,
                          " · ".join([z for z in [w["_ortjahr"], w["_masse"],
                                                  str(w["Exemplar"] or "")] if z]),
                          F_REG, 7, tw, 7 * 1.3, maxlines=1, color=(0.35, 0.35, 0.35))
        cy = draw_wrapped(c, x + pad, cy, w["_technik"], F_REG, 7, tw, 7 * 1.3,
                          maxlines=1, color=(0.35, 0.35, 0.35))
        cy = draw_wrapped(c, x + pad, cy, "Rahmen: " + str(w["Rahmenmodell"] or "—"),
                          F_REG, 7, tw, 7 * 1.3, maxlines=1, color=(0.35, 0.35, 0.35))

        kontakt = teilnehmende.get(str(w["Künstler:in"]), "")
        c.setFont(F_REG, 6.5)
        c.setFillColorRGB(0.55, 0.55, 0.55)
        c.drawString(x + pad, y + pad - 1 * mm,
                     (kontakt + " · " if kontakt else "") + AUSSTELLUNG["mail"])
        vers = f"VW {int(w['_vers'])} €" if w["_vers"] else ""
        c.drawRightString(x + bw - pad, y + pad - 1 * mm, vers)
    c.save()
    return (len(werke) + per - 1) // per


# ----------------------------------------------------------------- 05 Instagram-Wand
def pdf_instagram_wand(pfad, insta=None):
    """Aushang mit den Instagram-QR-Codes aller Teilnehmenden."""
    if T is None:
        print("   (übersprungen — teilnehmende.py nicht gefunden)")
        return 0
    insta = insta or {}
    namen = list(T.NAMEN)
    eintraege = []
    fehlt = []
    for n in namen:
        url = insta.get(n) or (T.INSTAGRAM.get(n) or {}).get("url", "")
        if url:
            handle = (T.INSTAGRAM.get(n) or {}).get("handle") or url.rstrip("/").rsplit("/", 1)[-1]
            eintraege.append((n, handle, url))
        else:
            fehlt.append(n)

    TXT, SEK, TER = FARBEN["text"], FARBEN["sekundaer"], FARBEN["tertiaer"]

    pw, ph = A4
    ml, mr, mt, mb = 18 * mm, 18 * mm, 26 * mm, 18 * mm
    cols, rows = 3, 4                     # nur 12 pro Seite -> viel Abstand
    per = cols * rows
    cw = (pw - ml - mr) / cols
    chh = (ph - mt - mb) / rows
    # QR bewusst klein halten, damit rundum genug Weißraum bleibt und
    # nicht versehentlich der Nachbarcode mitgescannt wird
    qs = min(cw - 22 * mm, chh - 26 * mm, 34 * mm)

    c = canvas.Canvas(pfad, pagesize=A4)
    c.setTitle("FOTOTAGE MUSTERSTADT 2026 — Instagram-Wand")
    if eintraege:
        zeichne_deckblatt(
            c, "Folg allen, die hier hängen", "Instagram-Wand", "",
        )
        c.showPage()
    for i, (name, handle, url) in enumerate(eintraege):
        if i % per == 0:
            if i:
                c.showPage()
            c.setFillColorRGB(*TXT)
            c.setFont(F_BOLD, 17)
            c.drawString(ml, ph - mt + 9 * mm, AUSSTELLUNG["titel"])
            c.setFont(F_REG, 10)
            c.setFillColorRGB(*SEK)
            c.drawString(ml, ph - mt + 3 * mm,
                         "Die Teilnehmenden auf Instagram — Code scannen und folgen")
        k = i % per
        cellx = ml + (k % cols) * cw
        celly = ph - mt - (k // cols) * chh
        midx = cellx + cw / 2
        qy = celly - qs - 5 * mm
        qr_zeichnen(c, qr_png(url, scale=8), midx - qs / 2, qy, qs)
        c.setFont(F_BOLD, 8.5)
        c.setFillColorRGB(*TXT)
        c.drawCentredString(midx, qy - 6 * mm, _kurz(name, F_BOLD, 8.5, cw - 6 * mm))
        c.setFont(F_REG, 7.5)
        c.setFillColorRGB(*TER)
        c.drawCentredString(midx, qy - 10.5 * mm, _kurz("@" + handle, F_REG, 7.5, cw - 6 * mm))
    if eintraege:
        c.setFont(F_REG, 7)
        c.setFillColorRGB(*TER)
        c.drawString(ml, mb - 8 * mm, f"{AUSSTELLUNG['web']} · {AUSSTELLUNG['insta']}")
        c.save()
    if fehlt:
        print(f"   ohne Instagram-Link ({len(fehlt)}): {', '.join(fehlt)}")
    return (len(eintraege) + per - 1) // per


def _kurz(txt, font, size, maxw):
    t = str(txt)
    while pdfmetrics.stringWidth(t, font, size) > maxw and len(t) > 3:
        t = t[:-2] + "…"
    return t


# ----------------------------------------------------------------- 04 Versicherungsliste
def pdf_versicherung(werke, pfad):
    pw, ph = A4
    ml, mr, mt, mb = 16 * mm, 16 * mm, 22 * mm, 20 * mm
    c = canvas.Canvas(pfad, pagesize=A4)
    c.setTitle("FOTOTAGE MUSTERSTADT 2026 — Versicherungs- und Kommissionsliste")

    zeichne_deckblatt(
        c, "Versicherungsliste", "Versicherungs- und Kommissionsliste",
        "Alle Arbeiten mit Preis und Versicherungswert, dazu Status und Käufer:in. "
        "Grundlage für die Versicherungsmeldung an Räumchen e.V. und für die Abrechnung "
        "nach der Ausstellung.",
        intern=True,
    )
    c.showPage()

    spalten = [("Nr", 12 * mm, "l"), ("Künstler:in", 34 * mm, "l"), ("Titel", 46 * mm, "l"),
               ("Maße", 22 * mm, "l"), ("Status", 26 * mm, "l"),
               ("Preis", 19 * mm, "r"), ("Vers.-Wert", 19 * mm, "r")]
    page = [0]

    def kopfzeile(y):
        c.setFont(F_BOLD, 7.5)
        c.setFillColorRGB(0.25, 0.25, 0.25)
        x = ml
        for name, br, al in spalten:
            (c.drawRightString(x + br - 2 * mm, y, name) if al == "r"
             else c.drawString(x, y, name))
            x += br
        c.setStrokeColorRGB(0.2, 0.2, 0.2)
        c.setLineWidth(0.6)
        c.line(ml, y - 2.5 * mm, pw - mr, y - 2.5 * mm)
        return y - 7 * mm

    def seitenkopf():
        page[0] += 1
        if page[0] == 1:
            c.setFont(F_BOLD, 15)
            c.setFillColorRGB(0, 0, 0)
            c.drawString(ml, ph - mt, "Versicherungs- und Kommissionsliste")
            c.setFont(F_REG, 9)
            c.setFillColorRGB(0.35, 0.35, 0.35)
            c.drawString(ml, ph - mt - 5 * mm,
                         f"{AUSSTELLUNG['titel']} · {AUSSTELLUNG['laufzeit']} · "
                         f"{AUSSTELLUNG['ort']} · Stand {date.today().strftime('%d.%m.%Y')}")
            c.drawString(ml, ph - mt - 9.5 * mm, "INTERN — nicht auslegen")
            return kopfzeile(ph - mt - 17 * mm)
        c.setFont(F_REG, 7.5)
        c.setFillColorRGB(0.55, 0.55, 0.55)
        c.drawString(ml, ph - mt, "Versicherungs- und Kommissionsliste · INTERN")
        return kopfzeile(ph - mt - 7 * mm)

    y = seitenkopf()
    summe_p = summe_v = 0
    for w in sorted(werke, key=lambda z: z["_nr"]):
        if y < mb + 14 * mm:
            c.setFont(F_REG, 7)
            c.setFillColorRGB(0.55, 0.55, 0.55)
            c.drawRightString(pw - mr, mb - 6 * mm, f"Seite {page[0]}")
            c.showPage()
            y = seitenkopf()
        x = ml
        werte = [
            w["_nr"], str(w["Künstler:in"]), str(w["Titel"]), w["_masse"], str(w["Status"]),
            (f"{int(w['_preis'])} €" if w["_preis"] else "—"),
            (f"{int(w['_vers'])} €" if w["_vers"] else "—"),
        ]
        c.setFont(F_REG, 8)
        c.setFillColorRGB(0.1, 0.1, 0.1)
        for (name, br, al), v in zip(spalten, werte):
            txt = v
            while pdfmetrics.stringWidth(txt, F_REG, 8) > br - 3 * mm and len(txt) > 3:
                txt = txt[:-2] + "…"
            (c.drawRightString(x + br - 2 * mm, y, txt) if al == "r"
             else c.drawString(x, y, txt))
            x += br
        summe_p += w["_preis"] or 0
        summe_v += w["_vers"] or 0
        c.setStrokeColorRGB(0.90, 0.90, 0.90)
        c.setLineWidth(0.25)
        c.line(ml, y - 2 * mm, pw - mr, y - 2 * mm)
        y -= 6 * mm

    y -= 2 * mm
    c.setStrokeColorRGB(0.2, 0.2, 0.2)
    c.setLineWidth(0.8)
    c.line(ml, y + 3 * mm, pw - mr, y + 3 * mm)
    c.setFont(F_BOLD, 8.5)
    c.setFillColorRGB(0, 0, 0)
    c.drawString(ml, y - 2 * mm, f"{len(werke)} Werke")
    x = ml + sum(br for _, br, _ in spalten[:5])
    c.drawRightString(x + spalten[5][1] - 2 * mm, y - 2 * mm, f"{int(summe_p)} €")
    c.drawRightString(x + spalten[5][1] + spalten[6][1] - 2 * mm, y - 2 * mm, f"{int(summe_v)} €")
    c.setFont(F_REG, 7)
    c.setFillColorRGB(0.55, 0.55, 0.55)
    c.drawString(ml, y - 9 * mm,
                 "Versicherungssumme an Räumchen e.V. melden. Bei Verkauf Zeile ergänzen: "
                 "Datum, Käufer:in, Zahlungsweg, Provisionsanteil.")
    c.drawRightString(pw - mr, mb - 6 * mm, f"Seite {page[0]}")
    c.save()
    return page[0]


# ----------------------------------------------------------------- CSV
def csv_export(werke, pfad):
    felder = ["Nr", "Wand", "Position", "Kuenstlerin", "Titel", "OrtLandJahr", "Technik",
              "Masse", "Blattmass", "Auflage", "Exemplar", "Rahmenmodell", "Passepartout",
              "Preis", "RahmenImPreis", "Status", "Preistext"]
    with open(pfad, "w", newline="", encoding="utf-8-sig") as f:
        wr = csv.writer(f, delimiter=";")
        wr.writerow(felder)
        for w in werke:
            wr.writerow([
                w["_nr"], w["Wand"], w["Position"], w["Künstler:in"], w["Titel"], w["_ortjahr"],
                w["_technik"], w["_masse"], w["_blattmass"], w["_auflage"], w["Exemplar"],
                w["Rahmenmodell"], w["Passepartout"],
                ("" if w["_preis"] is None else int(w["_preis"])),
                w["Rahmen im Preis"], w["Status"], preis_text(w),
            ])


# ----------------------------------------------------------------- Adobe-Exporte
# Adobe-taugliche Spaltennamen: keine Umlaute, keine Leerzeichen, keine Doppelpunkte.
ADOBE_FELDER = [
    ("Nr",          lambda w: w["_nr"]),
    ("Kuenstlerin", lambda w: str(w["Künstler:in"] or "")),
    ("Titel",       lambda w: str(w["Titel"] or "")),
    ("OrtJahr",     lambda w: w["_ortjahr"]),
    ("Technik",     lambda w: w["_technik"]),
    ("Masse",       lambda w: w["_masse"]),
    ("Auflage",     lambda w: w["_auflage"]),
    ("Hinweis",     lambda w: ("unverkäuflich" if w["_status"] == "nicht verkäuflich"
                               else "Preis siehe Werkliste")),
    ("Preistext",   lambda w: preis_text(w)),
    ("Wand",        lambda w: str(w["Wand"] or "")),
    ("Position",    lambda w: str(w["Position"] or "")),
]


def _adobe_zeilen(werke):
    return [[fn(w) for _n, fn in ADOBE_FELDER] for w in werke]


def export_indesign(werke, pfad):
    """Tabgetrennte Textdatei in UTF-16 LE — das von InDesign zuverlässigste Format
    für deutsche Umlaute in der Datenzusammenführung."""
    kopf = [n for n, _ in ADOBE_FELDER]
    zeilen = [kopf] + _adobe_zeilen(werke)
    text = "\r\n".join("\t".join(str(z).replace("\t", " ") for z in row) for row in zeilen)
    with open(pfad, "wb") as f:
        f.write("\ufeff".encode("utf-16-le"))       # BOM
        f.write(text.encode("utf-16-le"))


def export_illustrator_csv(werke, pfad):
    """Kommagetrennte CSV in UTF-8 ohne BOM — Format für das Variablenbedienfeld."""
    kopf = [n for n, _ in ADOBE_FELDER]
    with open(pfad, "w", newline="", encoding="utf-8") as f:
        wr = csv.writer(f, delimiter=",", quoting=csv.QUOTE_MINIMAL, lineterminator="\r\n")
        wr.writerow(kopf)
        for row in _adobe_zeilen(werke):
            wr.writerow(row)


def _xml_escape(t):
    return (str(t).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;"))


def export_illustrator_xml(werke, pfad):
    """Variablenbibliothek als XML — Notfallweg, falls die CSV nicht angenommen wird."""
    felder = [n for n, _ in ADOBE_FELDER]
    out = ['<?xml version="1.0" encoding="UTF-8"?>',
           '<svg xmlns="http://www.w3.org/2000/svg" '
           'xmlns:v="http://ns.adobe.com/Variables/1.0/">',
           '<v:variableSets xmlns:v="http://ns.adobe.com/Variables/1.0/">',
           '<v:variableSet v:varSetName="binding1" locked="none">',
           '<v:variables>']
    for n in felder:
        out.append(f'<v:variable v:varName="{n}" v:trait="textcontent"/>')
    out.append('</v:variables>')
    out.append('<v:sampleDataSets>')
    for w, zeile in zip(werke, _adobe_zeilen(werke)):
        name = _xml_escape(w["_nr"] or w["Titel"])
        out.append(f'<v:sampleDataSet v:dataSetName="{name}">')
        for n, v in zip(felder, zeile):
            out.append(f'<{n}><p>{_xml_escape(v)}</p></{n}>')
        out.append('</v:sampleDataSet>')
    out += ['</v:sampleDataSets>', '</v:variableSet>', '</v:variableSets>', '</svg>']
    with open(pfad, "w", encoding="utf-8") as f:
        f.write("\n".join(out))


def lade_teilnehmende(pfad):
    try:
        wb = load_workbook(pfad, data_only=True)
        if "Teilnehmende" not in wb.sheetnames:
            return {}
        ws = wb["Teilnehmende"]
        out = {}
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[0]:
                out[str(row[0])] = str(row[1] or "")
        return out
    except Exception:
        return {}


def _insta_url(roh):
    """Handle, @handle oder ganze URL -> https://www.instagram.com/handle/."""
    s = str(roh or "").strip()
    if not s:
        return ""
    if s.startswith("http"):
        return s
    s = s.lstrip("@").strip("/")
    s = s.replace("instagram.com/", "").replace("www.", "")
    return f"https://www.instagram.com/{s}/" if s else ""


def lade_instagram(pfad):
    """{Künstler:in: Instagram-URL} — aus dem Master-Blatt „Teilnehmende"
    (Spalte „Instagram"), ergänzt aus qr-codes/instagram.csv."""
    out = {}
    try:
        wb = load_workbook(pfad, data_only=True)
        if "Teilnehmende" in wb.sheetnames:
            ws = wb["Teilnehmende"]
            kopf = {str(c.value).strip(): i for i, c in enumerate(ws[1]) if c.value}
            ig = kopf.get("Instagram")
            for row in ws.iter_rows(min_row=2, values_only=True):
                if row[0] and ig is not None and ig < len(row):
                    u = _insta_url(row[ig])
                    if u:
                        out[str(row[0]).strip()] = u
    except Exception:
        pass
    if T is not None:
        for name, info in T.INSTAGRAM.items():
            if info.get("url") and name not in out:
                out[name] = info["url"]
    return out


_QR_CACHE = {}


def qr_png(url, scale=10):
    """PNG-Bytes eines QR-Codes zur URL, oder None wenn segno fehlt."""
    if not url:
        return None
    if url not in _QR_CACHE:
        try:
            import segno
            b = BytesIO()
            segno.make(url, error="m").save(b, kind="png", scale=scale, border=0)
            _QR_CACHE[url] = b.getvalue()
        except Exception:
            _QR_CACHE[url] = None
    return _QR_CACHE[url]


def qr_zeichnen(c, png, x, y, groesse):
    if png:
        c.drawImage(ImageReader(BytesIO(png)), x, y, groesse, groesse,
                    preserveAspectRatio=True, mask="auto")


# ----------------------------------------------------------------- Deckblatt
def _gesperrt(c, x, y, text, font, size, extra):
    """Zeichnet Text mit fester Buchstabenlaufweite (setCharSpace gibt es in
    älteren reportlab-Versionen nicht)."""
    c.setFont(font, size)
    for ch in str(text):
        c.drawString(x, y, ch)
        x += pdfmetrics.stringWidth(ch, font, size) + extra


def zeichne_deckblatt(c, name, unterzeile, beschreibung, intern=False):
    """Zeichnet ein A4-Titelblatt (Hochformat) auf die AKTUELLE Seite. Der
    Aufrufer ruft danach c.showPage() und setzt bei Bedarf mit c.setPageSize()
    das Format des eigentlichen Inhalts.

    Das Blatt ist eine schlichte, vorzeigbare Titelseite: Ausstellungstitel,
    große Dokumentbezeichnung mit Unterzeile, eine kurze Beschreibung und ein
    Impressumsblock (Ausstellung, Laufzeit, Ort, Veranstalter, Web). Bewusst
    ohne Druck-, Bogen- oder Schnitthinweise — das Blatt wird mitgedruckt und
    kann so mit ausgehängt oder ausgelegt werden. Interne Dokumente tragen
    oben rechts einen dezenten Vermerk „Intern — nicht aushängen“."""
    pw, ph = A4
    ml = mr = 24 * mm
    tw = pw - ml - mr
    TXT, SEK, TER, LIN = (FARBEN["text"], FARBEN["sekundaer"],
                          FARBEN["tertiaer"], FARBEN["linie"])

    c.setDash()
    y = ph - 40 * mm
    c.setFillColorRGB(*TER)
    _gesperrt(c, ml, y, AUSSTELLUNG["titel"].upper(), F_BOLD, 11, 2)

    if intern:
        c.setFont(F_REG, 8.5)
        c.setFillColorRGB(*FARBEN["signal"])
        c.drawRightString(pw - mr, y, "Intern — nicht aushängen")

    y -= 6 * mm
    c.setStrokeColorRGB(*TXT)
    c.setLineWidth(1)
    c.line(ml, y, pw - mr, y)

    y -= 40 * mm
    c.setFillColorRGB(*TXT)
    for ln in wrap(name, F_BOLD, 34, tw):
        c.setFont(F_BOLD, 34)
        c.drawString(ml, y, ln)
        y -= 41
    if unterzeile:
        y -= 4 * mm
        c.setFont(F_REG, 14)
        c.setFillColorRGB(*SEK)
        c.drawString(ml, y, unterzeile)

    if beschreibung:
        y -= 20 * mm
        draw_wrapped(c, ml, y, beschreibung, F_REG, 10.5, tw, 10.5 * 1.6, color=SEK)

    # --- Impressumsblock unten -------------------------------------------
    by = 42 * mm
    c.setStrokeColorRGB(*LIN)
    c.setLineWidth(0.5)
    c.line(ml, by + 16 * mm, pw - mr, by + 16 * mm)
    c.setFont(F_REG, 9.5)
    c.setFillColorRGB(*SEK)
    c.drawString(ml, by + 9 * mm, AUSSTELLUNG["untertitel"])
    c.drawString(ml, by + 3.5 * mm, f"{AUSSTELLUNG['laufzeit']} · {AUSSTELLUNG['ort']}")
    c.setFont(F_REG, 8.5)
    c.setFillColorRGB(*TER)
    c.drawString(ml, by - 2.5 * mm, f"{AUSSTELLUNG['web']} · {AUSSTELLUNG['insta']}")

    c.setFillColorRGB(0, 0, 0)


# ----------------------------------------------------------------- main
def main():
    ap = argparse.ArgumentParser(description="FOTOTAGE MUSTERSTADT 2026 — Druckdaten-Generator")
    ap.add_argument("--xlsx", default=P.output("Werkdaten_MASTER.xlsx"))
    ap.add_argument("--out", default=P.output("ausgabe"))
    ap.add_argument("--schildformat", type=str.lower, choices=FORMATE, default="a8",
                    help="DIN-A-Format der Wandschilder, a0 … a10 (quer); Standard a8. "
                         "Anordnung und Schriftgrößen ergeben sich automatisch, "
                         "Feinjustierung in design.json")
    ap.add_argument("--gruppierung", choices=["ohne", "Wand", "Künstler:in"], default="ohne")
    ap.add_argument("--spalten", type=int, choices=[1, 2], default=2)
    ap.add_argument("--ohne-preishinweis", action="store_true")
    ap.add_argument("--statuspunkte", action="store_true",
                    help="verkauft/reserviert als farbigen Punkt aufs Schild drucken")
    ap.add_argument("--ohne-schnittmarken", action="store_true")
    ap.add_argument("--ohne-qr", action="store_true",
                    help="keine Instagram-QR-Codes auf die Wandschilder drucken")
    ap.add_argument("--design", default=None,
                    help="design.json des Designers (Schriften, Größen, Farben)")
    ap.add_argument("--adobe", action="store_true",
                    help="zusätzlich Datendateien für InDesign und Illustrator schreiben")
    ap.add_argument("--nur", choices=["alle", "wandschilder", "werkliste", "rueckseiten",
                                      "versicherung", "csv", "adobe", "instagram"], default="alle")
    ap.add_argument("--trotz-fehler", action="store_true",
                    help="auch bei unvollständigen Daten drucken")
    ap.add_argument("--mit-beispielen", action="store_true",
                    help="die 14 Beispielwerke immer mitdrucken, auch wenn schon echte Werke da sind")
    ap.add_argument("--einzelkarte", metavar="NR",
                    help="nur ein Wandschild (Werknummer, z. B. 070) im gewählten "
                         "--schildformat erzeugen, eine Karte pro Blatt in Kartengröße — "
                         "für sehr große Wandbilder; separate Datei, ändert die "
                         "übrigen Ausgaben nicht")
    ap.add_argument("--einzelkarte-a3", metavar="NR",
                    help="Kurzform für: --einzelkarte NR --schildformat a3")
    args = ap.parse_args()

    font = register_fonts()
    design = lade_design(args.design)
    setze_farben(design)
    os.makedirs(args.out, exist_ok=True)

    insta = {} if args.ohne_qr else lade_instagram(args.xlsx)

    if args.nur == "instagram":
        b = pdf_instagram_wand(os.path.join(args.out, "05_Instagram_Wand.pdf"), insta)
        if b:
            print(f"05_Instagram_Wand.pdf      {b} Seite(n) — Aushang mit QR-Codes")
        return

    werke_alle = lade_werke(args.xlsx)
    if not werke_alle:
        sys.exit("Keine Werke gefunden — Blatt 'Werke' ist leer.")
    echt = [w for w in werke_alle if not w["_beispiel"]]
    demo = [w for w in werke_alle if w["_beispiel"]]
    if args.mit_beispielen:
        werke = werke_alle
    elif echt:
        # echte Werke vorhanden -> Beispieldaten nie mit ausgeben
        werke = echt
        if demo:
            print(f"{len(demo)} Beispielwerk(e) übersprungen (echte Werke vorhanden). "
                  f"--mit-beispielen erzwingt sie mit.")
    else:
        # nur Beispieldaten vorhanden -> zum Testen weiter anzeigen
        werke = werke_alle

    # Zentrale Sortierung nach Werknummer (die import_werkmeldungen.py bereits
    # alphabetisch nach Künstler:in vergibt) — gilt für ALLE Ausgaben, auch
    # für die, die nicht selbst noch einmal sortieren (Wandschilder, Rück-
    # seitenetiketten, CSV, InDesign/Illustrator). Werke ohne Nr landen zuletzt.
    werke.sort(key=lambda w: (w["_nr"] or "zzz", str(w["Künstler:in"] or "")))
    teiln = lade_teilnehmende(args.xlsx)

    if args.einzelkarte_a3:
        args.einzelkarte, args.schildformat = args.einzelkarte_a3, "a3"
    if args.einzelkarte:
        nr = args.einzelkarte.strip().zfill(3)
        treffer = [w for w in werke_alle if w["_nr"] == nr]
        if not treffer:
            sys.exit(f"FEHLER: keine Werk-Nr „{nr}“ gefunden.")
        w = treffer[0]
        name = str(w["Künstler:in"]).strip().replace(" ", "")
        os.makedirs(args.out, exist_ok=True)
        fmt = args.schildformat
        pfad = os.path.join(args.out, f"01_Wandschild_{fmt.upper()}_Nr{nr}_{name}.pdf")
        hinweise = pdf_wandschild_einzel(w, pfad, fmt, preishinweis=not args.ohne_preishinweis,
                                         statuspunkt=args.statuspunkte, insta=insta)
        print(f"{os.path.basename(pfad)}   Nr {nr} „{w['Titel']}“ ({w['Künstler:in']}), "
              f"{fmt.upper()} quer")
        for h in hinweise:
            print(f"    {h}")
        return

    probleme = pruefe(werke)
    print("\nFOTOTAGE MUSTERSTADT 2026 — Generator")
    print(f"Schrift: {font}")
    hinweis_bsp = " (nur Beispieldaten — echte Werke fehlen noch)" if not echt else ""
    print(f"Werke eingelesen: {len(werke)}{hinweis_bsp}")
    if probleme:
        print(f"\n{len(probleme)} Datenproblem(e):")
        for p in probleme[:40]:
            print(p)
        if len(probleme) > 40:
            print(f"  … und {len(probleme) - 40} weitere")
        if not args.trotz_fehler:
            print("\nAbbruch. Erst die Tabelle vervollständigen, "
                  "oder mit --trotz-fehler trotzdem drucken.")
            sys.exit(1)
        print("\n--trotz-fehler gesetzt, drucke weiter.\n")
    else:
        print("Datenprüfung: keine Fehler.\n")

    o = args.out
    tun = args.nur

    if tun in ("alle", "wandschilder"):
        n, boegen, hinweise = pdf_wandschilder(
            werke, os.path.join(o, "01_Wandschilder.pdf"), args.schildformat,
            preishinweis=not args.ohne_preishinweis, statuspunkt=args.statuspunkte,
            marken=not args.ohne_schnittmarken, insta=insta)
        qn = sum(1 for w in werke if insta.get(str(w["Künstler:in"]).strip()))
        print(f"01_Wandschilder.pdf        {n} Schilder auf {boegen} Bogen "
              f"({args.schildformat.upper()})"
              + (f", QR auf {qn}" if insta else ", ohne QR"))
        for h in hinweise:
            print(f"    {h}")

    if tun in ("alle", "werkliste"):
        s = pdf_werkliste(werke, os.path.join(o, "02_Werkliste.pdf"),
                          spalten=args.spalten, gruppe=args.gruppierung)
        g = "fortlaufend" if args.gruppierung == "ohne" else f"nach {args.gruppierung}"
        print(f"02_Werkliste.pdf           {s} Seite(n), {g}")

    if tun in ("alle", "rueckseiten"):
        b = pdf_rueckseiten(werke, os.path.join(o, "03_Rueckseitenetiketten.pdf"), teiln)
        print(f"03_Rueckseitenetiketten.pdf {b} Bogen à 12 Etiketten")

    if tun in ("alle", "versicherung"):
        s = pdf_versicherung(werke, os.path.join(o, "04_Versicherungsliste.pdf"))
        print(f"04_Versicherungsliste.pdf  {s} Seite(n) — INTERN")

    if tun in ("alle", "csv"):
        csv_export(werke, os.path.join(o, "Werkdaten_Serienbrief.csv"))
        print("Werkdaten_Serienbrief.csv  Fallback für Word-Serienbrief")

    if tun in ("alle", "instagram"):
        b = pdf_instagram_wand(os.path.join(o, "05_Instagram_Wand.pdf"), insta)
        if b:
            print(f"05_Instagram_Wand.pdf      {b} Seite(n) — Aushang mit QR-Codes")

    if args.adobe or tun == "adobe":
        ad = os.path.join(o, "Adobe")
        os.makedirs(ad, exist_ok=True)
        export_indesign(werke, os.path.join(ad, "InDesign_Datenzusammenfuehrung.txt"))
        export_illustrator_csv(werke, os.path.join(ad, "Illustrator_Variablen.csv"))
        export_illustrator_xml(werke, os.path.join(ad, "Illustrator_Variablenbibliothek.xml"))
        print(f"Adobe/                     3 Datendateien für InDesign und Illustrator "
              f"({len(werke)} Datensätze)")

    print(f"\nFertig. Alles in {os.path.abspath(o)}\n")


if __name__ == "__main__":
    main()
