#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Bereitet saal_digital_wandbild_poster_fineart.csv auf:

  * zerlegt die verdichteten Spalten (Material, Druckverfahren, Format,
    "Weitere Optionen") in einzelne, saubere Spalten
  * kategorisiert (Materialgruppe, Verfahrensgruppe, Seitenverhältnis, DIN …)
  * schreibt eine flache Tabelle als CSV
  * schreibt eine Arbeitsmappe mit Katalog + Zusammenfassungen als XLSX

Aufruf:  ./.venv/bin/python saal_katalog_aufbereiten.py
"""

import csv
import os
import re
from collections import defaultdict

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

HIER = os.path.dirname(os.path.abspath(__file__))
QUELLE = os.path.join(HIER, "saal_digital_wandbild_poster_fineart.csv")
CSV_OUT = os.path.join(HIER, "saal_digital_katalog.csv")
XLSX_OUT = os.path.join(HIER, "saal_digital_katalog.xlsx")

FONT = "Arial"
HEAD = PatternFill("solid", fgColor="D9D9D9")
SUB = PatternFill("solid", fgColor="EDEDED")

ZAHLWORT = {"Zwei": 2, "Drei": 3, "Vier": 4, "Fuenf": 5,
            "Sechs": 6, "Sieben": 7, "Acht": 8, "Neun": 9}

SV = {
    "~2:3": "2:3", "~3:4": "3:4", "~3:5": "3:5", "~4:5": "4:5",
    "Quadratisch": "1:1", "Panorama": "Panorama",
    "diverse": "diverse", "diverse (alle Seitenverhaeltnisse)": "diverse",
    "Mehrteilig": "mehrteilig",
}

DIN = {
    "10,5x14,8": "A6", "14,8x21": "A5", "21x29,7": "A4", "21x30": "A4",
    "29,7x42": "A3", "42x59,4": "A2", "59,4x84,1": "A1",
}


_UMLAUT = [
    ("waehlbar", "wählbar"), ("moeglich", "möglich"), ("verfuegbar", "verfügbar"),
    ("gebuerstetes", "gebürstet"), ("gebuerstete", "gebürstet"), ("gebuerstet", "gebürstet"),
    ("Weiss", "Weiß"), ("weiss", "weiß"), ("Groessen", "Größen"), ("Groesse", "Größe"),
    ("Staerke", "Stärke"), ("Aufhaengung", "Aufhängung"),
    ("Holzaufhaenger", "Holzaufhänger"), ("aufhaenger", "aufhänger"),
]


def de(s):
    for a, b in _UMLAUT:
        s = s.replace(a, b)
    return s


def num(s):
    s = s.strip().replace(",", ".")
    try:
        f = float(s)
        return int(f) if f.is_integer() else f
    except ValueError:
        return None


# ---------------------------------------------------------------- Produktlinie
def split_produktlinie(pl):
    pl = de(pl)
    m = re.match(r"Mehrteilige Wandbilder - (\w+) Teile( gekachelt)?", pl)
    if m:
        return "Mehrteilige Wandbilder", ZAHLWORT.get(m.group(1), None), \
            ("gekachelt" if m.group(2) else "nebeneinander")
    return pl, 1, ""


# ---------------------------------------------------------------- Material
def split_material(mat, verfahren):
    marke = ""
    if "Fujifilm" in verfahren:
        marke = "Fujifilm"
    elif "Hahnemuehle" in verfahren or mat.startswith("Hahnemuehle"):
        marke = "Hahnemühle"

    if mat.startswith("Fotopapier"):
        ober = mat.split(" ", 1)[1] if " " in mat else ""
        ober = {"glaenzend": "glänzend", "matt": "matt", "Silk": "Silk"}.get(ober, ober)
        return "Fotopapier (Fujifilm)", "Fotopapier", ober, marke or "Fujifilm"

    if mat.startswith("Hahnemuehle"):
        return "FineArt-Papier (Hahnemühle)", mat.replace("Hahnemuehle", "Hahnemühle"), "", "Hahnemühle"

    if "gebuerstet" in mat:
        farbe = re.search(r"\((\w+)\)", mat)
        return "Aluminium-Verbund", "Alu-Verbund gebürstet", \
            ("gebürstet " + farbe.group(1) if farbe else "gebürstet"), ""

    if mat.startswith("Alu-Verbund") or mat == "Alu-Verbund":
        return "Aluminium-Verbund", "Alu-Verbund", "", ""

    if mat.startswith("Acryl"):
        ober = "hinterklebt" if "hinterklebt" in mat else ""
        return "Acrylglas", "Acrylglas", ober, ""

    if mat.startswith("Hartschaum"):
        return "Hartschaumplatte", "Hartschaumplatte", "", ""

    if mat.startswith("GalleryPrint"):
        return "Glas / Galerie", "GalleryPrint", "", ""

    if "Schattenfugenrahmen" in mat:
        return "Leinwand", "Fotoleinwand", "", ""
    if "Akustik" in mat or "Schallschutz" in mat:
        return "Leinwand", "Akustik-Leinwand", "", ""
    if "Leinwand" in mat or "Canvas" in mat:
        return "Leinwand", "Baumwoll-Leinwand", "", ""

    return "Sonstige", mat, "", marke


# ---------------------------------------------------------------- Druckverfahren
def split_verfahren(v):
    if "Fujifilm" in v:
        return "Fotobelichtung / Digitaldruck", "Fotografisch"
    if "Hahnemuehle" in v or "FineArt-Inkjet" in v:
        return "FineArt-Inkjetdruck", "Pigment / FineArt"
    if v.startswith("Direktdruck"):
        return v, "Direktdruck"
    if v.startswith("Latex"):
        return v.replace(" (Standard)", ""), "Latexdruck"
    if v == "Standard":
        return "GalleryPrint", "Fotografisch"
    return v, "Sonstige"


# ---------------------------------------------------------------- Format
def split_format(f):
    """-> (breite, hoehe, format_text, din, gesamtformat)"""
    if f.lower().startswith("diverse"):
        return None, None, "diverse", "", ""

    base = f.split("(")[0].strip()
    din = DIN.get(base, "")

    m = re.match(r"^(\d+)x(\d+(?:,\d+)?)x(\d+(?:,\d+)?)$", base)  # mehrteilig
    if m:
        b, h = num(m.group(2)), num(m.group(3))
        g = re.search(r"gesamt ([\dx,]+)", f)
        return b, h, f"{b}x{h}", din, (g.group(1) if g else "")

    parts = base.split("x")
    if len(parts) == 2:
        b, h = num(parts[0]), num(parts[1])
        if b is None or h is None:
            return None, None, base, din, ""
        return b, h, f"{b}x{h}", din, ""

    return None, None, base, din, ""


# ---------------------------------------------------------------- Weitere Optionen
def split_optionen(o):
    """-> (aufhaengung, rahmenfarben, staerke, form, hinweis)"""
    auf = rahmen = staerke = form = hinweis = ""

    mst = re.search(r"Sta?erke:\s*([^;]+)", o)
    if mst:
        staerke = mst.group(1).strip()
        o = o.replace(mst.group(0), "").strip(" ;")

    mf = re.search(r"Form:\s*([^;]+)", o)
    if mf:
        form = mf.group(1).strip()
        o = o.replace(mf.group(0), "").strip(" ;")

    if o.startswith("Rahmenfarbe"):
        rest = re.sub(r"^Rahmenfarbe\s+wa?ehlbar[:\s(]*", "", o).rstrip(") ")
        return "", de(rest), staerke, form, ""

    o = re.sub(r"^(Aufhaengung/Rahmen|Aufhaengung):\s*", "", o)

    if ";" in o:
        auf, hinweis = (p.strip() for p in o.split(";", 1))
    elif re.search(r",\s*weiter", o):
        auf, rest = re.split(r",\s*weiter", o, maxsplit=1)
        auf, hinweis = auf.strip(), "weiter" + rest.strip()
    else:
        auf = o.strip()

    auf = de(auf.replace("Schattenfuge weiss", "Schattenfuge weiß"))
    # "ohne Aufhängung" (ggf. mit Klammerzusatz)  ->  "ohne (…)"
    m = re.fullmatch(r"ohne Aufhängung\s*(\(.*\))?", auf)
    if m:
        auf = "ohne " + m.group(1) if m.group(1) else "ohne"
    return auf, rahmen, staerke, form, de(hinweis)


SPALTEN = [
    "Kategorie", "Produktlinie", "Teile", "Anordnung",
    "Materialgruppe", "Material", "Oberflaeche", "Marke",
    "Druckverfahren", "Verfahrensgruppe",
    "Ausrichtung", "Seitenverhaeltnis",
    "Breite_cm", "Hoehe_cm", "Format", "DIN", "Gesamtformat",
    "Aufhaengung", "Rahmenfarben", "Staerke", "Form", "Hinweis",
]


def transform():
    src = list(csv.DictReader(open(QUELLE, encoding="utf-8-sig"), delimiter=";"))
    out = []
    for r in src:
        pl, teile, anordnung = split_produktlinie(r["Produktlinie"])
        mgrp, mat, ober, marke = split_material(
            r["Material/Papier (Oberflaeche)"], r["Druckverfahren"])
        verf, vgrp = split_verfahren(r["Druckverfahren"])
        b, h, fmt, din, gesamt = split_format(r["Format (cm)"])
        auf, rahmen, staerke, form, hinweis = split_optionen(
            r["Weitere Optionen (Aufhaengung/Rahmen/Farbe/Staerke/Form)"])
        if not staerke and re.search(r"\b5\s?mm\b", r["Material/Papier (Oberflaeche)"]):
            staerke = "5 mm"
        out.append({
            "Kategorie": r["Kategorie"],
            "Produktlinie": pl, "Teile": teile, "Anordnung": anordnung,
            "Materialgruppe": mgrp, "Material": mat, "Oberflaeche": ober, "Marke": marke,
            "Druckverfahren": verf, "Verfahrensgruppe": vgrp,
            "Ausrichtung": r["Ausrichtung"],
            "Seitenverhaeltnis": SV.get(r["Formatgruppe (Seitenverhaeltnis)"],
                                        r["Formatgruppe (Seitenverhaeltnis)"]),
            "Breite_cm": b, "Hoehe_cm": h, "Format": fmt, "DIN": din, "Gesamtformat": gesamt,
            "Aufhaengung": auf, "Rahmenfarben": rahmen, "Staerke": staerke,
            "Form": form, "Hinweis": hinweis,
        })
    return out


# ---------------------------------------------------------------- Ausgabe
def schreibe_csv(rows):
    with open(CSV_OUT, "w", encoding="utf-8-sig", newline="") as fh:
        wr = csv.DictWriter(fh, fieldnames=SPALTEN, delimiter=";", lineterminator="\n")
        wr.writeheader()
        wr.writerows(rows)
    print("geschrieben:", CSV_OUT, f"({len(rows)} Zeilen)")


def _stil_kopf(ws, ncols, row=1):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = Font(name=FONT, size=9, bold=True)
        cell.fill = HEAD
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 26
    ws.freeze_panes = ws.cell(row=row + 1, column=1).coordinate


def _fmt_spanne(werte):
    paare = sorted({(b, h) for b, h in werte if b and h}, key=lambda p: (p[0] * p[1], p))
    if not paare:
        return "", ""
    return f"{paare[0][0]}x{paare[0][1]}", f"{paare[-1][0]}x{paare[-1][1]}"


def schreibe_xlsx(rows):
    wb = Workbook()

    # ---- Blatt 1: Zusammenfassung ------------------------------------
    ws = wb.active
    ws.title = "Zusammenfassung"
    ws.sheet_view.showGridLines = False
    kat = defaultdict(int)
    for r in rows:
        kat[r["Kategorie"]] += 1
    total_pl = len({(r["Kategorie"], r["Produktlinie"]) for r in rows})
    total_mat = len({r["Material"] for r in rows})
    total_fmt = len({r["Format"] for r in rows if r["Breite_cm"]})
    vgrp = defaultdict(int)
    for r in rows:
        vgrp[r["Verfahrensgruppe"]] += 1

    zeilen = [
        ("Saal Digital — Katalog Wandbild / Poster / FineArt", ""),
        ("", ""),
        ("Zeilen (Produkt × Material × Ausrichtung × Format)", len(rows)),
        ("Kategorien", ", ".join(f"{k} ({v})" for k, v in sorted(kat.items()))),
        ("Produktlinien", total_pl),
        ("verschiedene Materialien", total_mat),
        ("verschiedene Formate (cm)", total_fmt),
        ("", ""),
        ("Zeilen je Verfahrensgruppe", ""),
    ]
    for k, v in sorted(vgrp.items(), key=lambda x: -x[1]):
        zeilen.append((f"    {k}", v))

    for i, (a, b) in enumerate(zeilen, start=1):
        ws.cell(row=i, column=1, value=a).font = Font(name=FONT, size=10,
                                                      bold=(b == "" and a != ""))
        ws.cell(row=i, column=2, value=b).font = Font(name=FONT, size=10)
    ws.cell(row=1, column=1).font = Font(name=FONT, size=13, bold=True)
    ws.column_dimensions["A"].width = 48
    ws.column_dimensions["B"].width = 60

    # ---- Blatt 2: Produktlinien -------------------------------------
    wp = wb.create_sheet("Produktlinien")
    kopf = ["Kategorie", "Produktlinie", "Materialgruppen", "Materialien",
            "Druckverfahren", "Ausrichtungen", "Seitenverhältnisse",
            "Formate", "kleinstes", "größtes", "Aufhängung / Optionen"]
    wp.append(kopf)
    grp = defaultdict(list)
    for r in rows:
        grp[(r["Kategorie"], r["Produktlinie"])].append(r)
    for (k, pl), rs in sorted(grp.items()):
        mn, mx = _fmt_spanne([(r["Breite_cm"], r["Hoehe_cm"]) for r in rs])
        opt = sorted({r["Aufhaengung"] for r in rs if r["Aufhaengung"]}) \
            or sorted({r["Rahmenfarben"] for r in rs if r["Rahmenfarben"]})
        wp.append([
            k, pl,
            ", ".join(sorted({r["Materialgruppe"] for r in rs})),
            ", ".join(sorted({r["Material"] for r in rs})),
            ", ".join(sorted({r["Druckverfahren"] for r in rs})),
            ", ".join(sorted({r["Ausrichtung"] for r in rs})),
            ", ".join(sorted({r["Seitenverhaeltnis"] for r in rs})),
            len({r["Format"] for r in rs if r["Breite_cm"]}),
            mn, mx,
            " | ".join(opt)[:200],
        ])
    _stil_kopf(wp, len(kopf))
    for c, wdt in zip("ABCDEFGHIJK", (14, 30, 24, 46, 30, 22, 20, 8, 10, 10, 60)):
        wp.column_dimensions[c].width = wdt
    wp.auto_filter.ref = f"A1:{get_column_letter(len(kopf))}{wp.max_row}"

    # ---- Blatt 3: Materialien --------------------------------------
    wm = wb.create_sheet("Materialien")
    kopf = ["Materialgruppe", "Material", "Marke", "Oberflächen",
            "Druckverfahren", "Produktlinien", "Formate", "kleinstes", "größtes"]
    wm.append(kopf)
    grp = defaultdict(list)
    for r in rows:
        grp[(r["Materialgruppe"], r["Material"])].append(r)
    for (mg, mat), rs in sorted(grp.items()):
        mn, mx = _fmt_spanne([(r["Breite_cm"], r["Hoehe_cm"]) for r in rs])
        wm.append([
            mg, mat,
            ", ".join(sorted({r["Marke"] for r in rs if r["Marke"]})),
            ", ".join(sorted({r["Oberflaeche"] for r in rs if r["Oberflaeche"]})),
            ", ".join(sorted({r["Druckverfahren"] for r in rs})),
            ", ".join(sorted({r["Produktlinie"] for r in rs})),
            len({r["Format"] for r in rs if r["Breite_cm"]}),
            mn, mx,
        ])
    _stil_kopf(wm, len(kopf))
    for c, wdt in zip("ABCDEFGHI", (24, 26, 14, 22, 30, 40, 8, 10, 10)):
        wm.column_dimensions[c].width = wdt
    wm.auto_filter.ref = f"A1:{get_column_letter(len(kopf))}{wm.max_row}"

    # ---- Blatt 4: Katalog (alle Zeilen) ---------------------------
    wk = wb.create_sheet("Katalog")
    wk.append(SPALTEN)
    for r in rows:
        wk.append([r[s] for s in SPALTEN])
    _stil_kopf(wk, len(SPALTEN))
    breiten = {"Kategorie": 14, "Produktlinie": 26, "Materialgruppe": 22, "Material": 26,
               "Druckverfahren": 26, "Verfahrensgruppe": 16, "Ausrichtung": 12,
               "Seitenverhaeltnis": 15, "Aufhaengung": 40, "Rahmenfarben": 34,
               "Hinweis": 40, "Format": 10, "Gesamtformat": 12}
    for i, s in enumerate(SPALTEN, start=1):
        wk.column_dimensions[get_column_letter(i)].width = breiten.get(s, 10)
    wk.auto_filter.ref = f"A1:{get_column_letter(len(SPALTEN))}{wk.max_row}"

    wb.save(XLSX_OUT)
    print("geschrieben:", XLSX_OUT, "(Zusammenfassung, Produktlinien, Materialien, Katalog)")


if __name__ == "__main__":
    rows = transform()
    schreibe_csv(rows)
    schreibe_xlsx(rows)
